"""Render the Airline All-Real Evidence Showcase from committed evidence only.

This Airline-specific integration reads six explicit committed evidence files,
builds one immutable presentation model, and renders deterministic presentation
artifacts. It performs no package discovery, model call, runtime execution, or
evidence mutation.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import stat
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

MODULE_ID = "run_airline_all_real_evidence_showcase_v01"
SHOWCASE_STATUS_PASS = "PASS"
FAILURE_REASON = "airline_all_real_evidence_showcase_generation_failed"

PAGE_WIDTH = 960.0
PAGE_HEIGHT = 540.0
PPTX_WIDTH_INCHES = 13.333333
PPTX_HEIGHT_INCHES = 7.5
PREVIEW_WIDTH = 1600
PREVIEW_HEIGHT = 900

OUTPUT_NAMES = (
    "README.md",
    "claim_evidence_matrix_v01.json",
    "executive_one_pager_v01.pdf",
    "hedgehog_os_airline_all_real_showcase_v01.pdf",
    "hedgehog_os_airline_all_real_showcase_v01.pptx",
    "technical_appendix_v01.pdf",
)
ALL_DELIVERABLE_NAMES = (
    "README.md",
    "hedgehog_os_airline_all_real_showcase_v01.pptx",
    "hedgehog_os_airline_all_real_showcase_v01.pdf",
    "executive_one_pager_v01.pdf",
    "technical_appendix_v01.pdf",
    "claim_evidence_matrix_v01.json",
    "SHA256SUMS",
)

PREFLIGHT_PATH = "docs/airline_all_real_evidence_showcase_preflight_v01.md"
HUMAN_STORY_PATH = "docs/airline_all_real_full_stack_human_story_v01.md"
ANCHOR_PATH = "docs/airline_all_real_full_stack_crypto_anchor_v01.json"
REPLAY_REPORT_PATH = "docs/airline_all_real_full_stack_replay_report_v01.json"
GENERATION_AUDIT_PATH = (
    "docs/audit_reports/"
    "auditor_airline_all_real_full_stack_v01_generation_anchor_publication.log"
)
REPLAY_AUDIT_PATH = (
    "docs/audit_reports/"
    "auditor_airline_all_real_full_stack_v01_anchored_replay.log"
)

SOURCE_COMMITS = {
    PREFLIGHT_PATH: "63b6e06",
    HUMAN_STORY_PATH: "a562c47",
    ANCHOR_PATH: "a701743",
    REPLAY_REPORT_PATH: "ec50c1f",
    GENERATION_AUDIT_PATH: "a701743",
    REPLAY_AUDIT_PATH: "ec50c1f",
}

ACTOR_IDS = (
    "tri_party_airline_orchestrator_llm",
    "tri_party_airline_semantic_architect_llm",
    "client_purchase_intent_reviewer_llm",
    "client_profile_privacy_reviewer_llm",
    "airline_offer_policy_reviewer_llm",
    "airline_fare_rules_vertical_cell_llm",
    "airline_seat_baggage_vertical_cell_llm",
    "airline_ticketing_policy_reviewer_llm",
    "bank_payment_policy_reviewer_llm",
    "bank_idempotency_risk_vertical_cell_llm",
    "bank_payment_status_explainer_llm",
    "tri_party_evidence_consistency_reviewer_llm",
)

HUMAN_ACTOR_LABELS = (
    "Tri-Party Orchestrator",
    "Semantic Architect",
    "Client Purchase Intent",
    "Client Privacy",
    "Airline Offer Policy",
    "Fare Rules Cell",
    "Seat & Baggage Cell",
    "Ticketing Policy",
    "Bank Payment Policy",
    "Idempotency Risk Cell",
    "Payment Status Explainer",
    "Cross-Root Consistency",
)

TRANSACTION_EVIDENCE_IDS = (
    "airline_hold_commit_packet:semantic_causal:001",
    "client_purchase_intent:client_001:001",
    "bank_payment_authorization_ref:mock_bank_a:001",
    "airline_ticket_issue_commit_packet:mock_airline_al:001",
    "mock_ticket_receipt:mock_airline_al:001",
    "mock_purchase_receipt:client_001:001",
)

ROOT_FINAL_IDS = (
    "client_root_final:001",
    "airline_root_final:001",
    "bank_root_final:001",
)

SHOWCASE_ENTRY_PATHS = (
    HUMAN_STORY_PATH,
    ANCHOR_PATH,
    REPLAY_REPORT_PATH,
    GENERATION_AUDIT_PATH,
    REPLAY_AUDIT_PATH,
    "docs/showcase/airline_all_real_full_stack_v01/claim_evidence_matrix_v01.json",
)

SLIDE_TITLES = (
    "Cover - central thesis",
    "Executive proof in 90 seconds",
    "The user's mock travel task and accepted result",
    "Complete architecture from user task to anchored Replay",
    "The twelve real Gemini actors and bounded topology",
    "Orchestration and Client-side actor evidence",
    "Airline-side actor evidence",
    "Bank-side and cross-root actor evidence",
    "Why the deterministic system selected Offer A",
    "Three sovereign Roots and authority separation",
    "Five-phase Ticket/Purchase Corridor and mock evidence card",
    "Transaction Artifact Ledger - 19 / 29 / 3",
    "Crypto generation - 9 / 11 and SELF_CONSISTENT_UNANCHORED",
    "Committed external Anchor and 19-row Replay PASS",
    "First honest FAIL_CLOSED run and distinct successful recovery",
    "What was proved, non-claims, and evidence entry points",
)

APPENDIX_TITLES = (
    "Evidence identity and commit lineage",
    *ACTOR_IDS,
    "Ledger timeline rows 0-9",
    "Ledger timeline rows 10-18",
    "Crypto, hashes, and committed Anchor",
    "Replay and zero-rerun boundary",
    "Claim-evidence reference",
    "Non-claims and verification instructions",
)

COLORS = {
    "background": "#07131F",
    "primary_panel": "#0E2233",
    "secondary_panel": "#122C3D",
    "primary_text": "#F4F7FA",
    "muted_text": "#A9BAC8",
    "cyan": "#29D3E2",
    "client": "#4C9FFF",
    "airline": "#17C3A2",
    "bank": "#A78BFA",
    "pass": "#43D17C",
    "advisory": "#FFB74D",
    "fail": "#FF6675",
    "border": "#28465D",
}

EXPECTED_MANIFEST_HASH = (
    "4f6d9abe351e986abf147150a5dd4e76816430cf80bb718b50fa3536719a2472"
)
EXPECTED_PACKAGE_HASH = (
    "00662506fc3afe95d619a5003b572c9c1c9475e5ba26f7759b694115d092fdfa"
)
EXPECTED_CHAIN_TAIL = (
    "1f56be26e253f878e8d165d06a4da4430c1c4ec2e22daea1fc98c57b2f3b40fe"
)
EXPECTED_ANCHOR_SHA256 = (
    "54ee8e6dcceaa846ecc4bf8968b0ab26c9bcb849663a76c0cb9d01d564c694b0"
)
EXPECTED_REPLAY_SHA256 = (
    "7ceaa5400ef6b35fb1e39f2bdc739a68cb9bb4373d82e7586feee901ef414b57"
)


@dataclass(frozen=True)
class EvidenceSource:
    path: str
    source_commit: str
    source_sha256: str
    text: str
    parsed_json: object | None


@dataclass(frozen=True)
class Claim:
    claim_id: str
    display_label: str
    observed_value: object
    observed_type: str
    evidence_availability: str
    source_path: str
    source_field_or_section: str
    source_sha256: str
    source_commit: str
    slide_numbers: tuple[int, ...]
    appendix_pages: tuple[int, ...]
    public_link_available: bool
    notes: str


@dataclass(frozen=True)
class ActorCard:
    actor_id: str
    human_role: str
    side: str
    bounded_information: str
    accepted_canonical_meaning: str
    structured_fields: str
    validation_result: str
    downstream_consumer: str
    authority_boundary: str


@dataclass(frozen=True)
class TimelineRow:
    replay_index: int
    ledger_index: int
    artifact_type: str
    artifact_id: str
    root_owner: str
    dependency_count: int
    is_root_final: bool


@dataclass(frozen=True)
class Metric:
    label: str
    value: str
    color: str


@dataclass(frozen=True)
class TextBlock:
    x: float
    y: float
    width: float
    height: float
    text: str
    font_size: float
    color: str
    bold: bool = False
    align: str = "left"
    role: str = "body"


@dataclass(frozen=True)
class DiagramNode:
    node_id: str
    x: float
    y: float
    width: float
    height: float
    label: str
    fill: str
    stroke: str
    text_color: str
    font_size: float = 18.0
    bold: bool = False
    shape: str = "rounded"


@dataclass(frozen=True)
class DiagramEdge:
    source_x: float
    source_y: float
    target_x: float
    target_y: float
    color: str


@dataclass(frozen=True)
class SlideScene:
    slide_number: int
    title: str
    claim_ids: tuple[str, ...]
    blocks: tuple[TextBlock, ...]
    nodes: tuple[DiagramNode, ...]
    edges: tuple[DiagramEdge, ...]


@dataclass(frozen=True)
class AppendixPage:
    page_number: int
    title: str
    claim_ids: tuple[str, ...]
    blocks: tuple[TextBlock, ...]
    nodes: tuple[DiagramNode, ...]


@dataclass(frozen=True)
class OnePagerModel:
    blocks: tuple[TextBlock, ...]
    nodes: tuple[DiagramNode, ...]
    edges: tuple[DiagramEdge, ...]


@dataclass(frozen=True)
class ShowcaseModel:
    evidence_sources: tuple[EvidenceSource, ...]
    claims: tuple[Claim, ...]
    actors: tuple[ActorCard, ...]
    timeline: tuple[TimelineRow, ...]
    metrics: tuple[Metric, ...]
    slides: tuple[SlideScene, ...]
    appendix_pages: tuple[AppendixPage, ...]
    one_pager: OnePagerModel


def _stable_error() -> ValueError:
    error = ValueError(FAILURE_REASON)
    error.__cause__ = None
    error.__context__ = None
    return error


def _hex_to_rgb(value: str) -> tuple[int, int, int]:
    raw = value.removeprefix("#")
    return (int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def _ascii_text(value: str) -> str:
    replacements = {
        "\u2192": "->",
        "\u2014": "-",
        "\u2013": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2022": "-",
    }
    normalized = value
    for source, target in replacements.items():
        normalized = normalized.replace(source, target)
    return normalized


def _strict_json(raw: bytes) -> object:
    def pairs_hook(pairs: list[tuple[str, object]]) -> dict[str, object]:
        result: dict[str, object] = {}
        for key, value in pairs:
            if key in result:
                raise ValueError("duplicate_key")
            result[key] = value
        return result

    def reject_constant(_: str) -> object:
        raise ValueError("invalid_constant")

    if raw.startswith(b"\xef\xbb\xbf"):
        raise ValueError("bom")
    return json.loads(
        raw.decode("utf-8"),
        object_pairs_hook=pairs_hook,
        parse_constant=reject_constant,
    )


def _read_source(repository_root: Path, relative_path: str) -> EvidenceSource:
    path = repository_root / relative_path
    path_stat = path.lstat()
    if stat.S_ISLNK(path_stat.st_mode) or not stat.S_ISREG(path_stat.st_mode):
        raise ValueError("source_type")
    raw = path.read_bytes()
    parsed: object | None = None
    if relative_path.endswith(".json"):
        parsed = _strict_json(raw)
    return EvidenceSource(
        path=relative_path,
        source_commit=SOURCE_COMMITS[relative_path],
        source_sha256=hashlib.sha256(raw).hexdigest(),
        text=raw.decode("utf-8"),
        parsed_json=parsed,
    )


def _source_map(sources: tuple[EvidenceSource, ...]) -> dict[str, EvidenceSource]:
    return {source.path: source for source in sources}


def _metadata_value(markdown: str, field_name: str) -> str:
    match = re.search(
        rf"^- {re.escape(field_name)}: (.+)$",
        markdown,
        flags=re.MULTILINE,
    )
    if match is None:
        raise ValueError("metadata")
    return match.group(1).strip()


def _actor_cards(markdown: str) -> tuple[ActorCard, ...]:
    label_map = (
        ("Human role", "human_role"),
        ("Root side or advisory side", "side"),
        ("Bounded information available to the actor", "bounded_information"),
        ("Accepted canonical answer", "accepted_canonical_meaning"),
        ("Important accepted structured fields", "structured_fields"),
        ("Validation result", "validation_result"),
        ("Downstream consumer", "downstream_consumer"),
        ("Authority boundary", "authority_boundary"),
    )
    cards: list[ActorCard] = []
    for index, actor_id in enumerate(ACTOR_IDS, 1):
        heading = f"### 3.{index} `{actor_id}`"
        start = markdown.find(heading)
        if start < 0:
            raise ValueError("actor_heading")
        next_heading = (
            f"### 3.{index + 1} `{ACTOR_IDS[index]}`"
            if index < len(ACTOR_IDS)
            else "The five causal semantic actors were"
        )
        end = markdown.find(next_heading, start + len(heading))
        if end < 0:
            raise ValueError("actor_boundary")
        section = markdown[start:end]
        values: dict[str, str] = {}
        for label, key in label_map:
            match = re.search(
                rf"^- \*\*{re.escape(label)}:\*\* (.+)$",
                section,
                flags=re.MULTILINE,
            )
            if match is None:
                raise ValueError("actor_field")
            values[key] = match.group(1).strip()
        if "PASS" not in values["validation_result"]:
            raise ValueError("actor_validation")
        if "accepted: true" not in values["validation_result"]:
            raise ValueError("actor_acceptance")
        if "Advisory only" not in values["authority_boundary"]:
            raise ValueError("actor_authority")
        cards.append(ActorCard(actor_id=actor_id, **values))
    return tuple(cards)


def _human_timeline(markdown: str) -> tuple[TimelineRow, ...]:
    section_start = markdown.index("## 8. The 19-Artifact Ledger")
    section_end = markdown.index("## 9. The Crypto Seal")
    rows: list[TimelineRow] = []
    for line in markdown[section_start:section_end].splitlines():
        if re.match(r"^\| \d+ \|", line) is None:
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        rows.append(
            TimelineRow(
                replay_index=int(cells[0]),
                ledger_index=int(cells[0]),
                artifact_type=cells[1],
                artifact_id=cells[2].strip("`"),
                root_owner=cells[3].strip("`"),
                dependency_count=int(cells[4]),
                is_root_final=cells[5].startswith("yes"),
            )
        )
    if len(rows) != 19:
        raise ValueError("human_timeline")
    return tuple(rows)


def _validate_sources(
    sources: tuple[EvidenceSource, ...],
) -> tuple[tuple[ActorCard, ...], tuple[TimelineRow, ...]]:
    mapping = _source_map(sources)
    human = mapping[HUMAN_STORY_PATH].text
    anchor = mapping[ANCHOR_PATH].parsed_json
    replay = mapping[REPLAY_REPORT_PATH].parsed_json
    generation_audit = mapping[GENERATION_AUDIT_PATH].text
    replay_audit = mapping[REPLAY_AUDIT_PATH].text
    preflight = mapping[PREFLIGHT_PATH].text
    if not isinstance(anchor, dict) or not isinstance(replay, dict):
        raise ValueError("json_root")

    expected_metadata = {
        "document_status": "HUMAN_EXPLANATION",
        "official_generation_status": "PASS",
        "provider_mode": "real_provider",
        "model": "gemini-2.5-flash",
        "stored_crypto_status": "SELF_CONSISTENT_UNANCHORED",
        "fresh_anchored_verification_status": "PASS",
        "replay_status": "PASS",
        "real_world_effects_count": "0",
    }
    for field_name, expected in expected_metadata.items():
        if _metadata_value(human, field_name) != expected:
            raise ValueError("human_metadata")
    actors = _actor_cards(human)
    human_rows = _human_timeline(human)
    for marker in (
        "## 15. Evidence Map",
        "| Claim | Observed value | Primary evidence | Field or section | Commit or hash |",
        "## 14. What This Demonstration Does Not Prove",
        "Gemini output is not truth",
        "Replay is not effect authorization",
    ):
        if marker not in human:
            raise ValueError("human_contract")

    anchor_expected = {
        "expected_manifest_core_hash": EXPECTED_MANIFEST_HASH,
        "source_package_hash": EXPECTED_PACKAGE_HASH,
        "chain_tail_hash": EXPECTED_CHAIN_TAIL,
        "package_generation_base_head": "3301ce3",
        "verification_status_at_publication": "SELF_CONSISTENT_UNANCHORED",
        "signature_mode": "UNSIGNED_PLACEHOLDER",
        "signature_verified": False,
        "anchored_pass_claimed": False,
    }
    for field_name, expected in anchor_expected.items():
        if anchor.get(field_name) != expected:
            raise ValueError("anchor_contract")
    if mapping[ANCHOR_PATH].source_sha256 != EXPECTED_ANCHOR_SHA256:
        raise ValueError("anchor_hash")

    replay_expected = {
        "replay_status": "PASS",
        "stored_verification_status": "SELF_CONSISTENT_UNANCHORED",
        "fresh_anchored_verification_status": "PASS",
        "external_anchor_supplied": True,
        "external_anchor_verified": True,
        "source_file_count": 9,
        "critical_package_file_count": 11,
        "ledger_entry_count": 19,
        "dependency_edge_count": 29,
        "root_final_count": 3,
        "timeline_row_count": 19,
        "provider_call_count": 0,
        "network_call_count": 0,
        "gemini_call_count": 0,
        "transaction_rerun_count": 0,
        "semantic_rerun_count": 0,
        "corridor_rerun_count": 0,
        "ledger_recollection_count": 0,
        "crypto_collection_count": 0,
        "replay_created_authority_count": 0,
        "replay_created_permission_count": 0,
        "replay_created_action_count": 0,
        "replay_created_packet_count": 0,
        "replay_created_receipt_count": 0,
        "replay_created_final_output_count": 0,
        "real_world_effects_count": 0,
        "source_bytes_unchanged": True,
        "critical_package_bytes_unchanged": True,
        "signature_verified": False,
    }
    for field_name, expected in replay_expected.items():
        if replay.get(field_name) != expected:
            raise ValueError("replay_contract")
    for field_name in (
        "integrity_verified",
        "continuity_verified",
        "ledger_verified",
        "manifest_binding_verified",
        "artifact_hashes_verified",
        "chain_order_verified",
        "dependency_graph_verified",
        "root_ownership_verified",
        "authority_evidence_boundaries_verified",
        "packet_lineage_verified",
        "receipt_lineage_verified",
        "transaction_identity_verified",
        "source_refs_verified",
        "secret_boundary_verified",
        "timeline_complete",
    ):
        if replay.get(field_name) is not True:
            raise ValueError("replay_integrity")
    replay_rows_raw = replay.get("reconstructed_timeline")
    if not isinstance(replay_rows_raw, list) or len(replay_rows_raw) != 19:
        raise ValueError("replay_timeline")
    replay_rows = tuple(
        TimelineRow(
            replay_index=row["replay_index"],
            ledger_index=row["ledger_index"],
            artifact_type=row["artifact_type"],
            artifact_id=row["artifact_id"],
            root_owner=row["root_owner"],
            dependency_count=row["dependency_count"],
            is_root_final=row["is_root_final"],
        )
        for row in replay_rows_raw
    )
    if replay_rows != human_rows:
        raise ValueError("timeline_mismatch")
    if mapping[REPLAY_REPORT_PATH].source_sha256 != EXPECTED_REPLAY_SHA256:
        raise ValueError("replay_hash")

    for marker in (
        "audit_status: PASS",
        "final status: PASS",
        "real provider calls: 12",
        "network calls: 12",
        "Gemini calls: 12",
        "actor validations PASS: 12",
        "actor validations FAIL: 0",
        "selected offer: offer:mock_airline_al:PAR-LIM:001",
        "Corridor execution count: 1",
        "Corridor status: PASS",
        "entries: 19",
        "dependency edges: 29",
        "Root finals: 3",
        "exact source files: 9",
        "exact critical files: 11",
        "stored Verification status: SELF_CONSISTENT_UNANCHORED",
        "earlier package promoted: false",
        "real-world effects: 0",
    ):
        if marker not in generation_audit:
            raise ValueError("generation_audit")
    for marker in (
        "audit_status: PASS",
        "Replay runner invoked during this audit: false",
        "verification status: SELF_CONSISTENT_UNANCHORED",
        "verification status: PASS",
        "Replay status: PASS",
        "timeline rows: 19",
        "accepted package changed during audit: false",
        "historical failed package changed during this audit: false",
        "historical failed package promoted: false",
        "accepted recovery package is distinct: true",
        "provider_call_count: 0",
        "network_call_count: 0",
        "gemini_call_count: 0",
        "real_world_effects_count: 0",
    ):
        if marker not in replay_audit:
            raise ValueError("replay_audit")
    for marker in (
        "preflight_status: READY_FOR_REVIEW",
        "main-deck slides and PDF pages: exactly 16",
        "generated deliverables: exactly 7",
        "airline_all_real_evidence_showcase_v01_slice_s2_renderer_and_artifacts",
    ):
        if marker not in preflight:
            raise ValueError("preflight_contract")
    return actors, replay_rows


def _observed_type(value: object) -> str:
    if type(value) is bool:
        return "boolean"
    if type(value) is int:
        return "integer"
    if type(value) is str:
        return "string"
    raise ValueError("observed_type")


def _build_claims(sources: tuple[EvidenceSource, ...]) -> tuple[Claim, ...]:
    mapping = _source_map(sources)
    slide_map = {
        "C01": (1, 3, 15),
        "C02": (2, 4, 5),
        "C03": (5, 6, 7, 8, 9),
        "C04": (3, 6),
        "C05": (3, 6, 7, 8, 9, 11),
        "C06": (5, 6, 7, 9),
        "C07": (1, 2, 4, 5, 8, 10),
        "C08": (4, 9, 11),
        "C09": (2, 11),
        "C10": (2, 4, 12),
        "C11": (2, 12),
        "C12": (2, 10, 12),
        "C13": (2, 13),
        "C14": (2, 13),
        "C15": (13, 14),
        "C16": (1, 2, 4, 14),
        "C17": (1, 4, 14, 15, 16),
        "C18": (2, 4, 14),
        "C19": (14, 16),
        "C20": (1, 2, 3, 8, 11, 16),
        "C21": (12, 14, 16),
        "C22": (15,),
        "C23": (15,),
        "C24": (10, 13, 14, 16),
    }
    appendix_map = {
        "C01": (1, 18),
        "C02": (1, 2, 3, 18),
        "C03": (*range(2, 14), 18),
        "C04": (1, 4, 18),
        "C05": (4, 6, 7, 8, 13, 18),
        "C06": (1, 2, 3, 18),
        "C07": (1, 13, 14, 15, 18),
        "C08": (1, 18),
        "C09": (1, 18),
        "C10": (14, 15, 18),
        "C11": (14, 15, 18),
        "C12": (14, 15, 18),
        "C13": (16, 18),
        "C14": (16, 18),
        "C15": (16, 18),
        "C16": (1, 16, 18),
        "C17": (1, 17, 18),
        "C18": (14, 15, 17, 18),
        "C19": (17, 18),
        "C20": (17, 18, 19),
        "C21": (16, 17, 18),
        "C22": (1, 18),
        "C23": (1, 18),
        "C24": (16, 17, 18, 19),
    }
    specs = (
        ("C01", "generation_pass", "PASS", GENERATION_AUDIT_PATH, "[ACCEPTED RECOVERY GENERATION]", "committed_audit_of_local_package"),
        ("C02", "real_gemini_calls_12", 12, GENERATION_AUDIT_PATH, "Gemini calls", "committed_audit_of_local_package"),
        ("C03", "actor_validations_12_pass_0_fail", "12 PASS / 0 FAIL", GENERATION_AUDIT_PATH, "actor validations", "committed_audit_of_local_package"),
        ("C04", "transaction_route_and_date", "tri_airline_purchase:PAR-LIM:2026-08-12:client_001 / PAR to LIM / 2026-08-12", HUMAN_STORY_PATH, "The User's Travel Task", "committed_repository_evidence"),
        ("C05", "offer_a_selected", "offer:mock_airline_al:PAR-LIM:001", GENERATION_AUDIT_PATH, "selected offer", "committed_audit_of_local_package"),
        ("C06", "four_bsep_projections_pass", "4 PASS", HUMAN_STORY_PATH, "How the System Chose Offer A", "committed_repository_evidence"),
        ("C07", "three_sovereign_roots", "ClientRoot / AirlineRoot / BankRoot", HUMAN_STORY_PATH, "Three Sovereign Roots", "committed_repository_evidence"),
        ("C08", "corridor_pass", "PASS", GENERATION_AUDIT_PATH, "Corridor status", "committed_audit_of_local_package"),
        ("C09", "corridor_execution_count_1", 1, GENERATION_AUDIT_PATH, "Corridor execution count", "committed_audit_of_local_package"),
        ("C10", "ledger_entries_19", 19, GENERATION_AUDIT_PATH, "entries", "committed_audit_of_local_package"),
        ("C11", "ledger_dependency_edges_29", 29, GENERATION_AUDIT_PATH, "dependency edges", "committed_audit_of_local_package"),
        ("C12", "ledger_root_finals_3", 3, GENERATION_AUDIT_PATH, "Root finals", "committed_audit_of_local_package"),
        ("C13", "crypto_source_files_9", 9, GENERATION_AUDIT_PATH, "exact source files", "committed_audit_of_local_package"),
        ("C14", "critical_package_files_11", 11, GENERATION_AUDIT_PATH, "exact critical files", "committed_audit_of_local_package"),
        ("C15", "stored_verification_unanchored", "SELF_CONSISTENT_UNANCHORED", ANCHOR_PATH, "verification_status_at_publication", "committed_repository_evidence"),
        ("C16", "committed_external_anchor", True, ANCHOR_PATH, "anchor_active_only_when_committed", "committed_repository_evidence"),
        ("C17", "fresh_anchored_verification_pass", "PASS", REPLAY_REPORT_PATH, "fresh_anchored_verification_status", "committed_repository_evidence"),
        ("C18", "replay_timeline_rows_19", 19, REPLAY_REPORT_PATH, "timeline_row_count", "committed_repository_evidence"),
        ("C19", "replay_provider_network_gemini_zero", "provider 0 / network 0 / Gemini 0", REPLAY_REPORT_PATH, "provider_call_count / network_call_count / gemini_call_count", "committed_repository_evidence"),
        ("C20", "real_world_effects_zero", 0, REPLAY_REPORT_PATH, "real_world_effects_count", "committed_repository_evidence"),
        ("C21", "package_bytes_unchanged", True, REPLAY_REPORT_PATH, "critical_package_bytes_unchanged", "committed_repository_evidence"),
        ("C22", "historical_failure_evidence_preserved", True, REPLAY_AUDIT_PATH, "[PRIOR FAIL-CLOSED HISTORY]", "committed_audit_of_local_package"),
        ("C23", "failed_package_not_promoted_and_recovery_distinct", "failed package promoted false; recovery package distinct true", REPLAY_AUDIT_PATH, "[PRIOR FAIL-CLOSED HISTORY]", "committed_audit_of_local_package"),
        ("C24", "signature_verified_false", False, REPLAY_REPORT_PATH, "signature_verified", "committed_repository_evidence"),
    )
    claims: list[Claim] = []
    for claim_id, label, value, source_path, field, availability in specs:
        source = mapping[source_path]
        claims.append(
            Claim(
                claim_id=claim_id,
                display_label=label,
                observed_value=value,
                observed_type=_observed_type(value),
                evidence_availability=availability,
                source_path=source_path,
                source_field_or_section=field,
                source_sha256=source.source_sha256,
                source_commit=source.source_commit,
                slide_numbers=slide_map[claim_id],
                appendix_pages=tuple(sorted(set(appendix_map[claim_id]))),
                public_link_available=True,
                notes="Accepted committed evidence.",
            )
        )
    return tuple(claims)


def _wrap_line(text: str, width: float, font_size: float, bold: bool) -> list[str]:
    font_name = "Helvetica-Bold" if bold else "Helvetica"
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if stringWidth(candidate, font_name, font_size) <= width:
            current = candidate
            continue
        if current:
            lines.append(current)
            current = ""
        if stringWidth(word, font_name, font_size) <= width:
            current = word
            continue
        chunk = ""
        for character in word:
            candidate_chunk = chunk + character
            if stringWidth(candidate_chunk, font_name, font_size) > width and chunk:
                lines.append(chunk)
                chunk = character
            else:
                chunk = candidate_chunk
        current = chunk
    if current:
        lines.append(current)
    return lines


def _wrap_text(text: str, width: float, font_size: float, bold: bool = False) -> str:
    wrapped: list[str] = []
    for paragraph in _ascii_text(text).splitlines():
        wrapped.extend(_wrap_line(paragraph, width, font_size, bold))
    return "\n".join(wrapped)


def _block(
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    font_size: float = 18.0,
    color: str = COLORS["primary_text"],
    bold: bool = False,
    align: str = "left",
    role: str = "body",
) -> TextBlock:
    return TextBlock(
        x=x,
        y=y,
        width=width,
        height=height,
        text=_wrap_text(text, width, font_size, bold),
        font_size=font_size,
        color=color,
        bold=bold,
        align=align,
        role=role,
    )


def _node(
    node_id: str,
    x: float,
    y: float,
    width: float,
    height: float,
    label: str,
    fill: str = COLORS["primary_panel"],
    stroke: str = COLORS["border"],
    text_color: str = COLORS["primary_text"],
    font_size: float = 18.0,
    bold: bool = False,
    shape: str = "rounded",
) -> DiagramNode:
    return DiagramNode(
        node_id=node_id,
        x=x,
        y=y,
        width=width,
        height=height,
        label=_wrap_text(label, width - 18.0, font_size, bold),
        fill=fill,
        stroke=stroke,
        text_color=text_color,
        font_size=font_size,
        bold=bold,
        shape=shape,
    )


def _edge(source: DiagramNode, target: DiagramNode, color: str = COLORS["cyan"]) -> DiagramEdge:
    return DiagramEdge(
        source_x=source.x + source.width,
        source_y=source.y + source.height / 2.0,
        target_x=target.x,
        target_y=target.y + target.height / 2.0,
        color=color,
    )


def _scene(
    number: int,
    title: str,
    claims: tuple[str, ...],
    blocks: tuple[TextBlock, ...] = (),
    nodes: tuple[DiagramNode, ...] = (),
    edges: tuple[DiagramEdge, ...] = (),
) -> SlideScene:
    fixed_blocks = (
        _block(28, 10, 260, 18, "HEDGEHOG OS", 11, COLORS["cyan"], True, role="branding"),
        _block(28, 32, 904, 60, title, 24, COLORS["primary_text"], True, role="title"),
        _block(
            28,
            512,
            760,
            16,
            f"Evidence: {', '.join(claims)}",
            10,
            COLORS["muted_text"],
            role="footer",
        ),
        _block(874, 512, 58, 16, str(number), 10, COLORS["muted_text"], False, "right", "slide_number"),
    )
    return SlideScene(number, title, claims, fixed_blocks + blocks, nodes, edges)


def _actor_panel(actor: ActorCard, index: int, x: float, y: float) -> DiagramNode:
    role = actor.human_role.split(".")[0]
    meaning = actor.accepted_canonical_meaning.split(";")[0]
    label = (
        f"{index}. {actor.actor_id}\n"
        f"{role}\n"
        f"{meaning}\n"
        "PASS / accepted true / ADVISORY ONLY"
    )
    color = COLORS["advisory"]
    if "client" in actor.actor_id:
        color = COLORS["client"]
    elif "airline" in actor.actor_id:
        color = COLORS["airline"]
    elif "bank" in actor.actor_id:
        color = COLORS["bank"]
    return _node(
        f"actor_{index}",
        x,
        y,
        438,
        190,
        label,
        COLORS["primary_panel"],
        color,
        COLORS["primary_text"],
        18,
    )


def _build_slides(
    actors: tuple[ActorCard, ...],
    timeline: tuple[TimelineRow, ...],
) -> tuple[SlideScene, ...]:
    slides: list[SlideScene] = []

    chain_labels = ("Gemini evidence", "Root decisions", "Corridor", "Ledger", "Anchor", "Replay PASS")
    chain_nodes = tuple(
        _node(f"cover_{i}", 45 + i * 150, 250, 125, 68, label, COLORS["primary_panel"], COLORS["cyan"], font_size=18, bold=True)
        for i, label in enumerate(chain_labels)
    )
    slides.append(
        _scene(
            1,
            SLIDE_TITLES[0],
            ("C01", "C07", "C17", "C20"),
            blocks=(
                _block(
                    70,
                    92,
                    820,
                    118,
                    "HEDGEHOG OS separates LLM reasoning from authority, then preserves the accepted decision path as a deterministic, cryptographically anchored, replayable evidence chain.",
                    25,
                    COLORS["primary_text"],
                    True,
                    "center",
                ),
            ),
            nodes=chain_nodes
            + (
                _node("badge_12", 185, 370, 135, 55, "12 Gemini", COLORS["secondary_panel"], COLORS["advisory"], font_size=18, bold=True),
                _node("badge_3", 340, 370, 135, 55, "3 Roots", COLORS["secondary_panel"], COLORS["client"], font_size=18, bold=True),
                _node("badge_pass", 495, 370, 135, 55, "Replay PASS", COLORS["secondary_panel"], COLORS["pass"], font_size=18, bold=True),
                _node("badge_zero", 650, 370, 135, 55, "0 effects", COLORS["secondary_panel"], COLORS["cyan"], font_size=18, bold=True),
            ),
            edges=tuple(_edge(chain_nodes[i], chain_nodes[i + 1]) for i in range(5)),
        )
    )

    metric_specs = (
        ("12", "real Gemini actors", COLORS["advisory"]),
        ("3", "sovereign Roots", COLORS["client"]),
        ("1", "deterministic transaction", COLORS["cyan"]),
        ("19", "Ledger artifacts", COLORS["airline"]),
        ("29", "dependency edges", COLORS["airline"]),
        ("3", "Root finals", COLORS["client"]),
        ("9", "Crypto source files", COLORS["cyan"]),
        ("11", "critical package files", COLORS["cyan"]),
        ("1", "committed Anchor", COLORS["pass"]),
        ("19", "Replay rows", COLORS["pass"]),
        ("0", "real-world effects", COLORS["pass"]),
    )
    metric_nodes = tuple(
        _node(
            f"metric_{i}",
            42 + (i % 4) * 225,
            102 + (i // 4) * 118,
            202,
            92,
            f"{value}\n{label}",
            COLORS["primary_panel"],
            color,
            font_size=20,
            bold=True,
        )
        for i, (value, label, color) in enumerate(metric_specs)
    )
    slides.append(
        _scene(
            2,
            SLIDE_TITLES[1],
            ("C02", "C07", "C09", "C10", "C11", "C12", "C13", "C14", "C16", "C18", "C20"),
            nodes=metric_nodes,
        )
    )

    slides.append(
        _scene(
            3,
            SLIDE_TITLES[2],
            ("C01", "C04", "C05", "C20"),
            nodes=(
                _node("route", 55, 110, 285, 210, "PAR -> LIM\n2026-08-12\nPreference A", COLORS["primary_panel"], COLORS["cyan"], font_size=27, bold=True),
                _node("offer", 365, 110, 540, 210, "Selected Offer A\noffer:mock_airline_al:PAR-LIM:001\ntri_airline_purchase:PAR-LIM:2026-08-12:client_001", COLORS["primary_panel"], COLORS["pass"], font_size=21, bold=True),
                _node("mock", 115, 345, 730, 110, "MOCK / PROOF-LEVEL EVIDENCE\nNOT A REAL TICKET - NOT A REAL BOOKING - NOT A REAL PAYMENT", COLORS["secondary_panel"], COLORS["advisory"], font_size=22, bold=True),
            ),
        )
    )

    architecture_labels = (
        "User task",
        "BSEP",
        "12 actors",
        "3 Roots",
        "5-phase\nCorridor",
        "19-entry Ledger",
        "Crypto Manifest",
        "Committed Anchor",
        "19-row Replay PASS",
    )
    architecture_nodes = tuple(
        _node(f"arch_{i}", 16 + i * 105, 230, 94, 96, label, COLORS["primary_panel"], COLORS["cyan"], font_size=18, bold=True)
        for i, label in enumerate(architecture_labels)
    )
    slides.append(
        _scene(
            4,
            SLIDE_TITLES[3],
            ("C02", "C07", "C08", "C10", "C16", "C17", "C18"),
            blocks=(
                _block(50, 105, 190, 52, "REASONING", 18, COLORS["advisory"], True, "center"),
                _block(255, 105, 190, 52, "AUTHORITY", 18, COLORS["client"], True, "center"),
                _block(460, 105, 190, 52, "DETERMINISTIC", 18, COLORS["airline"], True, "center"),
                _block(665, 105, 190, 52, "EVIDENCE + VERIFY", 18, COLORS["pass"], True, "center"),
            ),
            nodes=architecture_nodes,
            edges=tuple(_edge(architecture_nodes[i], architecture_nodes[i + 1]) for i in range(8)),
        )
    )

    topology_border_colors = (
        COLORS["advisory"],
        COLORS["advisory"],
        COLORS["client"],
        COLORS["client"],
        COLORS["airline"],
        COLORS["airline"],
        COLORS["airline"],
        COLORS["airline"],
        COLORS["bank"],
        COLORS["bank"],
        COLORS["bank"],
        COLORS["advisory"],
    )
    actor_nodes = tuple(
        _node(
            f"topology_{i}",
            35 + (i % 3) * 300,
            100 + (i // 3) * 96,
            275,
            76,
            HUMAN_ACTOR_LABELS[i],
            COLORS["primary_panel"],
            topology_border_colors[i],
            font_size=18,
            bold=True,
        )
        for i, _actor in enumerate(actors)
    )
    slides.append(
        _scene(
            5,
            SLIDE_TITLES[4],
            ("C02", "C03", "C06", "C07"),
            blocks=(
                _block(650, 474, 280, 34, "5 causal / 7 generic / all advisory", 18, COLORS["muted_text"], True, "right"),
            ),
            nodes=actor_nodes,
        )
    )

    for slide_number, actor_slice, title, claims in (
        (6, actors[0:4], SLIDE_TITLES[5], ("C03", "C04", "C05", "C06")),
        (7, actors[4:8], SLIDE_TITLES[6], ("C03", "C05", "C06")),
        (8, actors[8:12], SLIDE_TITLES[7], ("C03", "C05", "C07", "C20")),
    ):
        slides.append(
            _scene(
                slide_number,
                title,
                claims,
                nodes=tuple(
                    _actor_panel(
                        actor,
                        ACTOR_IDS.index(actor.actor_id) + 1,
                        32 + (i % 2) * 456,
                        92 + (i // 2) * 194,
                    )
                    for i, actor in enumerate(actor_slice)
                ),
            )
        )

    funnel_labels = (
        "Bounded constraints",
        "BSEP",
        "4 projections",
        "12 accepted actors",
        "Canonical evidence",
        "ClientRoot",
        "AirlineRoot",
        "Offer A",
    )
    funnel_nodes = tuple(
        _node(f"funnel_{i}", 35 + i * 112, 215, 96, 94, label, COLORS["primary_panel"], COLORS["cyan"], font_size=18, bold=i >= 5)
        for i, label in enumerate(funnel_labels)
    )
    slides.append(
        _scene(
            9,
            SLIDE_TITLES[8],
            ("C03", "C05", "C06", "C08"),
            blocks=(
                _block(80, 105, 800, 58, "No provider override - no silent fallback - provider output is not truth or authority", 21, COLORS["advisory"], True, "center"),
                _block(250, 365, 460, 62, "offer:mock_airline_al:PAR-LIM:001", 24, COLORS["pass"], True, "center"),
            ),
            nodes=funnel_nodes,
            edges=tuple(_edge(funnel_nodes[i], funnel_nodes[i + 1]) for i in range(7)),
        )
    )

    root_nodes = (
        _node("client_root", 45, 120, 270, 292, "ClientRoot\nclient constraints\npurchase intent\nclient acceptance\nClientRoot final", COLORS["primary_panel"], COLORS["client"], font_size=21, bold=True),
        _node("airline_root", 345, 120, 270, 292, "AirlineRoot\noffer resolution\nhold\nticket issue intent\nmock ticket receipt\nAirlineRoot final", COLORS["primary_panel"], COLORS["airline"], font_size=21, bold=True),
        _node("bank_root", 645, 120, 270, 292, "BankRoot\npayment authorization boundary\npayment evidence\nBankRoot final", COLORS["primary_panel"], COLORS["bank"], font_size=21, bold=True),
    )
    slides.append(
        _scene(
            10,
            SLIDE_TITLES[9],
            ("C07", "C12", "C24"),
            blocks=(
                _block(115, 445, 730, 40, "Gemini != Root   Reviewer != Root   Crypto != Root   Replay != Root", 20, COLORS["advisory"], True, "center"),
            ),
            nodes=root_nodes,
        )
    )

    phase_labels = (
        "1. Airline offer + hold",
        "2. Client purchase intent",
        "3. Bank authorization",
        "4. Airline ticket intent",
        "5. Client completion",
    )
    phase_nodes = tuple(
        _node(f"phase_{i}", 30 + i * 183, 95, 165, 72, label, COLORS["primary_panel"], COLORS["cyan"], font_size=18, bold=True)
        for i, label in enumerate(phase_labels)
    )
    slides.append(
        _scene(
            11,
            SLIDE_TITLES[10],
            ("C05", "C08", "C09", "C20"),
            blocks=(
                _block(35, 188, 205, 30, "Hold packet", 18, COLORS["muted_text"], True),
                _block(250, 188, 675, 30, TRANSACTION_EVIDENCE_IDS[0], 18),
                _block(35, 222, 205, 30, "Client purchase intent", 18, COLORS["muted_text"], True),
                _block(250, 222, 675, 30, TRANSACTION_EVIDENCE_IDS[1], 18),
                _block(35, 256, 205, 30, "Bank authorization ref", 18, COLORS["muted_text"], True),
                _block(250, 256, 675, 30, TRANSACTION_EVIDENCE_IDS[2], 18),
                _block(35, 290, 205, 30, "Ticket issue intent", 18, COLORS["muted_text"], True),
                _block(250, 290, 675, 30, TRANSACTION_EVIDENCE_IDS[3], 18),
                _block(35, 324, 205, 30, "Mock ticket receipt", 18, COLORS["muted_text"], True),
                _block(250, 324, 675, 30, TRANSACTION_EVIDENCE_IDS[4], 18),
                _block(35, 358, 205, 30, "Mock purchase receipt", 18, COLORS["muted_text"], True),
                _block(250, 358, 675, 30, TRANSACTION_EVIDENCE_IDS[5], 18),
                _block(35, 392, 205, 30, "Root finals", 18, COLORS["muted_text"], True),
                _block(250, 392, 675, 30, " | ".join(ROOT_FINAL_IDS), 18),
                _block(250, 438, 675, 42, "MOCK EVIDENCE - NO REAL TICKET / BOOKING / PAYMENT", 18, COLORS["advisory"], True, "center"),
            ),
            nodes=phase_nodes,
            edges=tuple(_edge(phase_nodes[i], phase_nodes[i + 1]) for i in range(4)),
        )
    )

    graph_nodes: list[DiagramNode] = []
    for row in timeline:
        column = row.replay_index % 5
        line = row.replay_index // 5
        fill = COLORS["secondary_panel"]
        stroke = COLORS["advisory"]
        if "client" in row.root_owner:
            stroke = COLORS["client"]
        elif "airline" in row.root_owner:
            stroke = COLORS["airline"]
        elif "bank" in row.root_owner:
            stroke = COLORS["bank"]
        if row.is_root_final:
            fill = stroke
        graph_nodes.append(
            _node(
                f"ledger_{row.replay_index}",
                40 + column * 105,
                105 + line * 90,
                58,
                58,
                str(row.replay_index),
                fill,
                stroke,
                COLORS["primary_text"],
                18,
                True,
                "circle",
            )
        )
    graph_edges = tuple(
        DiagramEdge(
            graph_nodes[i].x + graph_nodes[i].width,
            graph_nodes[i].y + graph_nodes[i].height / 2,
            graph_nodes[i + 1].x,
            graph_nodes[i + 1].y + graph_nodes[i + 1].height / 2,
            COLORS["border"],
        )
        for i in range(18)
    )
    slides.append(
        _scene(
            12,
            SLIDE_TITLES[11],
            ("C10", "C11", "C12", "C21"),
            blocks=(
                _block(590, 105, 310, 95, "19 artifacts\n29 dependency edges\n3 Root finals", 24, COLORS["pass"], True, "center"),
                _block(590, 220, 310, 140, "Blue: ClientRoot\nTeal: AirlineRoot\nViolet: BankRoot\nAmber: advisory/evidence", 18, COLORS["muted_text"]),
                _block(590, 390, 310, 55, "All dependencies point backward.", 18, COLORS["primary_text"], True, "center"),
            ),
            nodes=tuple(graph_nodes),
            edges=graph_edges,
        )
    )

    crypto_labels = ("9 source files", "19 hashes", "Ordered chain", "Manifest Core", "11 critical files")
    crypto_nodes = tuple(
        _node(f"crypto_{i}", 55 + i * 180, 190, 145, 88, label, COLORS["primary_panel"], COLORS["cyan"], font_size=19, bold=True)
        for i, label in enumerate(crypto_labels)
    )
    slides.append(
        _scene(
            13,
            SLIDE_TITLES[12],
            ("C13", "C14", "C15", "C24"),
            blocks=(
                _block(100, 100, 760, 54, "Generation status: SELF_CONSISTENT_UNANCHORED", 25, COLORS["advisory"], True, "center"),
                _block(100, 340, 760, 90, "No external Anchor at generation. No anchored PASS claimed.\nUNSIGNED_PLACEHOLDER / signature verified false.\nIntegrity is not semantic truth.", 20, COLORS["primary_text"], False, "center"),
            ),
            nodes=crypto_nodes,
            edges=tuple(_edge(crypto_nodes[i], crypto_nodes[i + 1]) for i in range(4)),
        )
    )

    slides.append(
        _scene(
            14,
            SLIDE_TITLES[13],
            ("C15", "C16", "C17", "C18", "C19", "C21", "C24"),
            nodes=(
                _node("moment_1_heading", 55, 115, 385, 108, "MOMENT 1\nStored package Verification", COLORS["primary_panel"], COLORS["advisory"], font_size=22, bold=True),
                _node("moment_1_status", 55, 238, 385, 58, "SELF_CONSISTENT_UNANCHORED", COLORS["secondary_panel"], COLORS["advisory"], font_size=18, bold=True),
                _node("moment_1_boundary", 55, 311, 385, 64, "Anchor absent\nSignature unverified", COLORS["primary_panel"], COLORS["advisory"], font_size=20, bold=True),
                _node("moment_2", 520, 115, 385, 260, "MOMENT 2\nCommitted external Anchor\nFresh Verification PASS\nAnchor supplied + verified\n19-row Replay PASS", COLORS["primary_panel"], COLORS["pass"], font_size=22, bold=True),
            ),
            blocks=(
                _block(125, 410, 710, 62, "Replay provider / network / Gemini: 0 / 0 / 0   Package bytes unchanged: true", 19, COLORS["cyan"], True, "center"),
            ),
            edges=(DiagramEdge(440, 245, 520, 245, COLORS["cyan"]),),
        )
    )

    slides.append(
        _scene(
            15,
            SLIDE_TITLES[14],
            ("C22", "C23", "C01", "C17"),
            nodes=(
                _node("failure", 45, 112, 400, 310, "FIRST OFFICIAL PACKAGE\n12 real Gemini calls completed\nLedger compatibility false positive\nFAIL_CLOSED\nNo Anchor / no Replay\nPreserved / not promoted", COLORS["primary_panel"], COLORS["fail"], font_size=21, bold=True),
                _node("recovery", 515, 112, 400, 310, "DISTINCT RECOVERY\nRepair commit 38ad0b0\n411 focused tests\n1230 compatibility tests\nSeparate package\nGeneration PASS\nAnchored Replay PASS", COLORS["primary_panel"], COLORS["pass"], font_size=21, bold=True),
            ),
            edges=(DiagramEdge(445, 265, 515, 265, COLORS["cyan"]),),
        )
    )

    slides.append(
        _scene(
            16,
            SLIDE_TITLES[15],
            ("C17", "C19", "C20", "C21", "C24"),
            nodes=(
                _node("proved", 45, 105, 420, 275, "WHAT WAS PROVED\nBounded real LLM reasoning\nRoot authority separation\nDeterministic Corridor\nCausal Ledger\nByte-integrity Crypto Seal\nCommitted Anchor\nZero-rerun Replay", COLORS["primary_panel"], COLORS["pass"], font_size=20, bold=True),
                _node("not_proved", 495, 105, 420, 275, "WHAT WAS NOT PROVED\nNo real ticket / booking / payment\nNo production integration\nNo PKI / signer authentication\nNo Root Attestation\nCrypto is not semantic truth\nReplay is not effect authorization", COLORS["primary_panel"], COLORS["advisory"], font_size=20, bold=True),
            ),
            blocks=(
                _block(
                    55,
                    395,
                    850,
                    100,
                    "Evidence entry points:\n"
                    f"{SHOWCASE_ENTRY_PATHS[0]} | {SHOWCASE_ENTRY_PATHS[1]}\n"
                    f"{SHOWCASE_ENTRY_PATHS[2]} | {SHOWCASE_ENTRY_PATHS[3]}\n"
                    f"{SHOWCASE_ENTRY_PATHS[4]} | {SHOWCASE_ENTRY_PATHS[5]}",
                    10,
                    COLORS["cyan"],
                    role="evidence_path",
                ),
            ),
        )
    )
    return tuple(slides)


def _appendix_page(
    number: int,
    title: str,
    claims: tuple[str, ...],
    blocks: tuple[TextBlock, ...],
    nodes: tuple[DiagramNode, ...] = (),
) -> AppendixPage:
    fixed = (
        _block(28, 12, 280, 18, "HEDGEHOG OS / TECHNICAL APPENDIX", 10, COLORS["cyan"], True, role="branding"),
        _block(28, 34, 880, 38, title, 25, COLORS["primary_text"], True, role="title"),
        _block(28, 512, 760, 16, f"Evidence: {', '.join(claims)}", 10, COLORS["muted_text"], role="footer"),
        _block(874, 512, 58, 16, str(number), 10, COLORS["muted_text"], False, "right", "page_number"),
    )
    return AppendixPage(number, title, claims, fixed + blocks, nodes)


def _build_appendix(
    actors: tuple[ActorCard, ...],
    timeline: tuple[TimelineRow, ...],
    claims: tuple[Claim, ...],
) -> tuple[AppendixPage, ...]:
    pages: list[AppendixPage] = [
        _appendix_page(
            1,
            APPENDIX_TITLES[0],
            ("C01", "C02", "C07", "C16", "C17"),
            (
                _block(55, 100, 850, 300, "Evidence base: 63b6e06\nHuman Story: a562c47\nAnchor: a701743\nReplay Report: ec50c1f\nGeneration audit: a701743\nReplay audit: ec50c1f\n\nGeneration PASS / committed Anchor / Replay PASS / zero effects", 21, COLORS["primary_text"], True),
            ),
        )
    ]
    for page_number, actor in enumerate(actors, 2):
        pages.append(
            _appendix_page(
                page_number,
                actor.actor_id,
                ("C02", "C03", "C05", "C07"),
                (
                    _block(45, 95, 870, 325, f"Human role: {actor.human_role}\n\nSide: {actor.side}\n\nBounded information: {actor.bounded_information}\n\nAccepted canonical meaning: {actor.accepted_canonical_meaning}\n\nStructured fields: {actor.structured_fields}", 15, COLORS["primary_text"]),
                    _block(45, 420, 870, 76, f"Validation: PASS / accepted true   Downstream: {actor.downstream_consumer}\nAuthority boundary: {actor.authority_boundary}", 14, COLORS["advisory"], True),
                ),
            )
        )

    for page_number, subset, title in (
        (14, timeline[:10], APPENDIX_TITLES[13]),
        (15, timeline[10:], APPENDIX_TITLES[14]),
    ):
        lines = [
            f"{row.replay_index:02d} | {row.artifact_type} | {row.artifact_id} | {row.root_owner} | deps {row.dependency_count} | {'ROOT FINAL' if row.is_root_final else 'trace'}"
            for row in subset
        ]
        pages.append(
            _appendix_page(
                page_number,
                title,
                ("C10", "C11", "C12", "C18"),
                (_block(35, 88, 890, 400, "\n".join(lines), 12.5, COLORS["primary_text"]),),
            )
        )

    pages.append(
        _appendix_page(
            16,
            APPENDIX_TITLES[15],
            ("C13", "C14", "C15", "C16", "C24"),
            (
                _block(45, 92, 870, 380, f"Manifest Core:\n{EXPECTED_MANIFEST_HASH}\n\nSource package:\n{EXPECTED_PACKAGE_HASH}\n\nChain tail:\n{EXPECTED_CHAIN_TAIL}\n\nAnchor SHA-256:\n{EXPECTED_ANCHOR_SHA256}\n\nReplay Report SHA-256:\n{EXPECTED_REPLAY_SHA256}\n\nStored status: SELF_CONSISTENT_UNANCHORED\nSignature verified: false", 14.5, COLORS["primary_text"]),
            ),
        )
    )
    pages.append(
        _appendix_page(
            17,
            APPENDIX_TITLES[16],
            ("C17", "C18", "C19", "C20", "C21", "C24"),
            (
                _block(45, 95, 870, 340, "Replay status: PASS\nTimeline rows: 19\nExternal Anchor supplied: true\nExternal Anchor verified: true\nPackage bytes unchanged: true\n\ntransaction reruns: 0\nsemantic reruns: 0\nCorridor reruns: 0\nLedger recollections: 0\nCrypto recollections: 0\nprovider / network / Gemini: 0 / 0 / 0\ncreated authority / permission / action / packet / receipt / FinalOutput: 0\nreal-world effects: 0", 18, COLORS["primary_text"]),
            ),
        )
    )
    claim_lines = [
        f"{claim.claim_id} | {claim.display_label} | {claim.observed_value} | {claim.source_path}"
        for claim in claims
    ]
    pages.append(
        _appendix_page(
            18,
            APPENDIX_TITLES[17],
            tuple(claim.claim_id for claim in claims),
            (_block(30, 86, 900, 410, "\n".join(claim_lines), 10.5, COLORS["primary_text"]),),
        )
    )
    pages.append(
        _appendix_page(
            19,
            APPENDIX_TITLES[18],
            ("C20", "C24"),
            (
                _block(55, 100, 850, 305, "NOT A REAL TICKET\nNOT A REAL BOOKING\nNOT A REAL PAYMENT\nNOT PRODUCTION\nGEMINI OUTPUT IS NOT TRUTH OR AUTHORITY\nCRYPTO VALIDITY IS NOT SEMANTIC TRUTH\nTHE COMMITTED HASH ANCHOR IS NOT SIGNER AUTHENTICATION\nNO PKI / NO ROOT ATTESTATION\nREPLAY IS NOT EFFECT AUTHORIZATION", 20, COLORS["advisory"], True, "center"),
                _block(55, 430, 850, 55, "Verify SHA256SUMS, inspect the claim matrix, and follow repository-relative committed evidence paths.", 16, COLORS["cyan"], True, "center"),
            ),
        )
    )
    return tuple(pages)


def _build_one_pager() -> OnePagerModel:
    width, height = landscape(A4)
    return OnePagerModel(
        blocks=(
            _block(28, 18, 220, 18, "HEDGEHOG OS", 12, COLORS["cyan"], True, role="branding"),
            _block(28, 44, width - 56, 42, "Airline All-Real Evidence Showcase v0.1", 25, COLORS["primary_text"], True, role="title"),
            _block(40, 95, width - 80, 58, "LLM reasoning is separated from authority; the accepted decision path is deterministic, anchored, and replayable.", 20, COLORS["primary_text"], True, "center"),
            _block(45, 172, width - 90, 38, "12 / 3 / 19 / 29 / 9 / 11 / 19 / 0", 28, COLORS["pass"], True, "center"),
            _block(45, 220, width - 90, 42, "PAR -> LIM | 2026-08-12 | Offer A | MOCK / PROOF-LEVEL EVIDENCE", 18, COLORS["advisory"], True, "center"),
            _block(45, 445, width - 90, 44, "Stored Verification: SELF_CONSISTENT_UNANCHORED -> committed Anchor -> fresh PASS -> Replay PASS\nZero Gemini reruns. First failure preserved; distinct recovery accepted.", 14, COLORS["primary_text"], True, "center"),
            _block(45, 490, width - 90, 28, "No real ticket, booking, payment, PKI, signer authentication, Root Attestation, or production claim.", 10, COLORS["muted_text"], True, "center"),
            _block(
                45,
                520,
                width - 90,
                66,
                "Evidence entry points:\n"
                f"{SHOWCASE_ENTRY_PATHS[0]} | {SHOWCASE_ENTRY_PATHS[1]}\n"
                f"{SHOWCASE_ENTRY_PATHS[2]} | {SHOWCASE_ENTRY_PATHS[3]}\n"
                f"{SHOWCASE_ENTRY_PATHS[4]} | {SHOWCASE_ENTRY_PATHS[5]}",
                10,
                COLORS["cyan"],
                role="evidence_path",
            ),
        ),
        nodes=(
            _node("client_one", 45, 290, 220, 142, "ClientRoot\nconstraints\npurchase intent\nclient final", COLORS["primary_panel"], COLORS["client"], font_size=18, bold=True),
            _node("airline_one", 310, 290, 220, 142, "AirlineRoot\noffer / hold\nticket intent\nairline final", COLORS["primary_panel"], COLORS["airline"], font_size=18, bold=True),
            _node("bank_one", 575, 290, 220, 142, "BankRoot\nauthorization evidence\nbank final", COLORS["primary_panel"], COLORS["bank"], font_size=18, bold=True),
        ),
        edges=(),
    )


def _validate_text_fit(text: str, width: float, height: float, font_size: float, bold: bool) -> None:
    font_name = "Helvetica-Bold" if bold else "Helvetica"
    line_height = font_size * 1.22
    lines = text.splitlines() or [""]
    if len(lines) * line_height > height + 0.01:
        raise ValueError("text_height")
    if any(stringWidth(line, font_name, font_size) > width + 0.01 for line in lines):
        raise ValueError("text_width")


def _rectangles_overlap(
    first: tuple[float, float, float, float],
    second: tuple[float, float, float, float],
) -> bool:
    ax, ay, aw, ah = first
    bx, by, bw, bh = second
    return ax < bx + bw and ax + aw > bx and ay < by + bh and ay + ah > by


def _validate_scene(scene: SlideScene) -> None:
    if not scene.nodes and len(scene.blocks) <= 4:
        raise ValueError("empty_slide")
    if sum(block.role == "title" for block in scene.blocks) != 1:
        raise ValueError("slide_title")
    if sum(block.role == "footer" for block in scene.blocks) != 1:
        raise ValueError("slide_footer")
    for block in scene.blocks:
        if min(block.x, block.y, block.width, block.height) < 0:
            raise ValueError("block_bounds")
        if block.x + block.width > PAGE_WIDTH or block.y + block.height > PAGE_HEIGHT:
            raise ValueError("block_bounds")
        if block.role == "body" and block.font_size < 18:
            raise ValueError("body_font")
        if block.role == "footer" and block.font_size < 10:
            raise ValueError("footer_font")
        _validate_text_fit(block.text, block.width, block.height, block.font_size, block.bold)
    node_rects: list[tuple[float, float, float, float]] = []
    for node in scene.nodes:
        if node.x < 0 or node.y < 0 or node.x + node.width > PAGE_WIDTH or node.y + node.height > PAGE_HEIGHT:
            raise ValueError("node_bounds")
        if node.font_size < 18:
            raise ValueError("node_font")
        _validate_text_fit(node.label, node.width - 18, node.height - 12, node.font_size, node.bold)
        rect = (node.x, node.y, node.width, node.height)
        if any(_rectangles_overlap(rect, existing) for existing in node_rects):
            raise ValueError("node_overlap")
        node_rects.append(rect)
    content_rects = [
        (block.x, block.y, block.width, block.height)
        for block in scene.blocks
        if block.role == "body"
    ]
    for index, rect in enumerate(content_rects):
        if any(_rectangles_overlap(rect, other) for other in content_rects[index + 1 :]):
            raise ValueError("block_overlap")
        if any(_rectangles_overlap(rect, node_rect) for node_rect in node_rects):
            raise ValueError("block_node_overlap")


def _validate_auxiliary_page(
    *,
    blocks: tuple[TextBlock, ...],
    nodes: tuple[DiagramNode, ...],
    page_width: float,
    page_height: float,
) -> None:
    rectangles: list[tuple[float, float, float, float]] = []
    for block in blocks:
        if min(block.x, block.y, block.width, block.height) < 0:
            raise ValueError("aux_block_bounds")
        if block.x + block.width > page_width or block.y + block.height > page_height:
            raise ValueError("aux_block_bounds")
        if block.font_size < 10:
            raise ValueError("aux_font")
        _validate_text_fit(block.text, block.width, block.height, block.font_size, block.bold)
        rect = (block.x, block.y, block.width, block.height)
        if any(_rectangles_overlap(rect, existing) for existing in rectangles):
            raise ValueError("aux_overlap")
        rectangles.append(rect)
    for node in nodes:
        if node.x < 0 or node.y < 0 or node.x + node.width > page_width or node.y + node.height > page_height:
            raise ValueError("aux_node_bounds")
        if node.font_size < 10:
            raise ValueError("aux_node_font")
        _validate_text_fit(node.label, node.width - 18, node.height - 12, node.font_size, node.bold)
        rect = (node.x, node.y, node.width, node.height)
        if any(_rectangles_overlap(rect, existing) for existing in rectangles):
            raise ValueError("aux_overlap")
        rectangles.append(rect)


def build_showcase_model_v01(repository_root: str | Path) -> ShowcaseModel:
    root = Path(repository_root)
    if not root.is_dir() or root.is_symlink():
        raise _stable_error()
    source_paths = (
        PREFLIGHT_PATH,
        HUMAN_STORY_PATH,
        ANCHOR_PATH,
        REPLAY_REPORT_PATH,
        GENERATION_AUDIT_PATH,
        REPLAY_AUDIT_PATH,
    )
    try:
        sources = tuple(_read_source(root, path) for path in source_paths)
        actors, timeline = _validate_sources(sources)
        claims = _build_claims(sources)
        metrics = (
            Metric("real Gemini actors", "12", COLORS["advisory"]),
            Metric("sovereign Roots", "3", COLORS["client"]),
            Metric("Ledger artifacts", "19", COLORS["airline"]),
            Metric("dependency edges", "29", COLORS["airline"]),
            Metric("source / critical files", "9 / 11", COLORS["cyan"]),
            Metric("Replay rows", "19", COLORS["pass"]),
            Metric("real-world effects", "0", COLORS["pass"]),
        )
        slides = _build_slides(actors, timeline)
        appendix = _build_appendix(actors, timeline, claims)
        one_pager = _build_one_pager()
        if len(slides) != 16 or len(appendix) != 19 or len(claims) != 24:
            raise ValueError("model_geometry")
        for scene in slides:
            _validate_scene(scene)
        for page in appendix:
            if sum(block.role == "title" for block in page.blocks) != 1:
                raise ValueError("appendix_title")
            if sum(block.role == "footer" for block in page.blocks) != 1:
                raise ValueError("appendix_footer")
            _validate_auxiliary_page(
                blocks=page.blocks,
                nodes=page.nodes,
                page_width=PAGE_WIDTH,
                page_height=PAGE_HEIGHT,
            )
        one_page_width, one_page_height = landscape(A4)
        _validate_auxiliary_page(
            blocks=one_pager.blocks,
            nodes=one_pager.nodes,
            page_width=one_page_width,
            page_height=one_page_height,
        )
        return ShowcaseModel(
            evidence_sources=sources,
            claims=claims,
            actors=actors,
            timeline=timeline,
            metrics=metrics,
            slides=slides,
            appendix_pages=appendix,
            one_pager=one_pager,
        )
    except Exception:
        raise _stable_error() from None


def _pptx_color(hex_value: str):
    from pptx.dml.color import RGBColor

    return RGBColor(*_hex_to_rgb(hex_value))


def _add_pptx_text(slide, block: TextBlock) -> None:
    box = slide.shapes.add_textbox(Pt(block.x), Pt(block.y), Pt(block.width), Pt(block.height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for line_index, line in enumerate(block.text.splitlines() or [""]):
        paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }[block.align]
        paragraph.space_after = 0
        run = paragraph.runs[0]
        run.font.name = "Arial"
        run.font.size = Pt(block.font_size)
        run.font.bold = block.bold
        run.font.color.rgb = _pptx_color(block.color)


def _add_pptx_node(slide, node: DiagramNode) -> None:
    shape_type = {
        "rounded": MSO_SHAPE.ROUNDED_RECTANGLE,
        "rect": MSO_SHAPE.RECTANGLE,
        "circle": MSO_SHAPE.OVAL,
    }[node.shape]
    shape = slide.shapes.add_shape(
        shape_type,
        Pt(node.x),
        Pt(node.y),
        Pt(node.width),
        Pt(node.height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _pptx_color(node.fill)
    shape.line.color.rgb = _pptx_color(node.stroke)
    shape.line.width = Pt(1.2)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.margin_left = Pt(7)
    frame.margin_right = Pt(7)
    frame.margin_top = Pt(5)
    frame.margin_bottom = Pt(5)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for line_index, line in enumerate(node.label.splitlines() or [""]):
        paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_after = 0
        run = paragraph.runs[0]
        run.font.name = "Arial"
        run.font.size = Pt(node.font_size)
        run.font.bold = node.bold
        run.font.color.rgb = _pptx_color(node.text_color)


def _add_pptx_edge(slide, edge: DiagramEdge) -> None:
    line = slide.shapes.add_connector(
        1,
        Pt(edge.source_x),
        Pt(edge.source_y),
        Pt(edge.target_x),
        Pt(edge.target_y),
    )
    line.line.color.rgb = _pptx_color(edge.color)
    line.line.width = Pt(1.5)
    size = 8.0
    triangle = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Pt(edge.target_x - size),
        Pt(edge.target_y - size / 2),
        Pt(size),
        Pt(size),
    )
    triangle.rotation = 90
    triangle.fill.solid()
    triangle.fill.fore_color.rgb = _pptx_color(edge.color)
    triangle.line.fill.background()


def _normalize_pptx(raw_path: Path, output_path: Path) -> None:
    with zipfile.ZipFile(raw_path, "r") as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(
        output_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as target:
        for filename, data in sorted(entries):
            info = zipfile.ZipInfo(filename, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            target.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)


def _render_pptx(model: ShowcaseModel, output_path: Path, work_dir: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(PPTX_WIDTH_INCHES)
    presentation.slide_height = Inches(PPTX_HEIGHT_INCHES)
    properties = presentation.core_properties
    properties.title = "Hedgehog OS Airline All-Real Evidence Showcase v0.1"
    properties.subject = "Artifact-backed proof-level Airline evidence showcase"
    properties.author = "Hedgehog OS"
    properties.last_modified_by = "Hedgehog OS"
    properties.created = datetime(2000, 1, 1)
    properties.modified = datetime(2000, 1, 1)
    properties.revision = 1
    blank_layout = presentation.slide_layouts[6]
    for scene in model.slides:
        slide = presentation.slides.add_slide(blank_layout)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = _pptx_color(COLORS["background"])
        for edge in scene.edges:
            _add_pptx_edge(slide, edge)
        for node in scene.nodes:
            _add_pptx_node(slide, node)
        for block in scene.blocks:
            _add_pptx_text(slide, block)
    raw_path = work_dir / "raw_showcase.pptx"
    presentation.save(raw_path)
    _normalize_pptx(raw_path, output_path)
    raw_path.unlink()


def _draw_pdf_text(pdf: canvas.Canvas, block: TextBlock, page_height: float) -> None:
    pdf.setFillColorRGB(*[component / 255 for component in _hex_to_rgb(block.color)])
    font_name = "Helvetica-Bold" if block.bold else "Helvetica"
    pdf.setFont(font_name, block.font_size)
    line_height = block.font_size * 1.22
    lines = block.text.splitlines() or [""]
    start_y = page_height - block.y - (block.height - len(lines) * line_height) / 2 - block.font_size
    for index, line in enumerate(lines):
        y = start_y - index * line_height
        if block.align == "center":
            pdf.drawCentredString(block.x + block.width / 2, y, line)
        elif block.align == "right":
            pdf.drawRightString(block.x + block.width, y, line)
        else:
            pdf.drawString(block.x, y, line)


def _draw_pdf_node(pdf: canvas.Canvas, node: DiagramNode, page_height: float) -> None:
    fill = tuple(component / 255 for component in _hex_to_rgb(node.fill))
    stroke = tuple(component / 255 for component in _hex_to_rgb(node.stroke))
    pdf.setFillColorRGB(*fill)
    pdf.setStrokeColorRGB(*stroke)
    pdf.setLineWidth(1.2)
    bottom = page_height - node.y - node.height
    if node.shape == "circle":
        pdf.ellipse(node.x, bottom, node.x + node.width, bottom + node.height, fill=1, stroke=1)
    elif node.shape == "rect":
        pdf.rect(node.x, bottom, node.width, node.height, fill=1, stroke=1)
    else:
        pdf.roundRect(node.x, bottom, node.width, node.height, 8, fill=1, stroke=1)
    _draw_pdf_text(
        pdf,
        TextBlock(
            node.x + 9,
            node.y + 6,
            node.width - 18,
            node.height - 12,
            node.label,
            node.font_size,
            node.text_color,
            node.bold,
            "center",
            "node",
        ),
        page_height,
    )


def _draw_pdf_edge(pdf: canvas.Canvas, edge: DiagramEdge, page_height: float) -> None:
    color = tuple(component / 255 for component in _hex_to_rgb(edge.color))
    pdf.setStrokeColorRGB(*color)
    pdf.setFillColorRGB(*color)
    source_y = page_height - edge.source_y
    target_y = page_height - edge.target_y
    pdf.setLineWidth(1.5)
    pdf.line(edge.source_x, source_y, edge.target_x, target_y)
    pdf.line(edge.target_x, target_y, edge.target_x - 7, target_y + 4)
    pdf.line(edge.target_x, target_y, edge.target_x - 7, target_y - 4)


def _new_pdf(path: Path, pagesize: tuple[float, float]) -> canvas.Canvas:
    pdf = canvas.Canvas(
        str(path),
        pagesize=pagesize,
        pageCompression=1,
        invariant=1,
    )
    pdf.setTitle("Hedgehog OS Airline All-Real Evidence Showcase v0.1")
    pdf.setSubject("Artifact-backed proof-level Airline evidence showcase")
    pdf.setAuthor("Hedgehog OS")
    pdf.setCreator("Hedgehog OS")
    return pdf


def _render_deck_pdf(model: ShowcaseModel, output_path: Path) -> None:
    pdf = _new_pdf(output_path, (PAGE_WIDTH, PAGE_HEIGHT))
    for scene in model.slides:
        pdf.setFillColorRGB(*[component / 255 for component in _hex_to_rgb(COLORS["background"])])
        pdf.rect(0, 0, PAGE_WIDTH, PAGE_HEIGHT, fill=1, stroke=0)
        for edge in scene.edges:
            _draw_pdf_edge(pdf, edge, PAGE_HEIGHT)
        for node in scene.nodes:
            _draw_pdf_node(pdf, node, PAGE_HEIGHT)
        for block in scene.blocks:
            _draw_pdf_text(pdf, block, PAGE_HEIGHT)
        pdf.showPage()
    pdf.save()


def _render_appendix_pdf(model: ShowcaseModel, output_path: Path) -> None:
    pdf = _new_pdf(output_path, (PAGE_WIDTH, PAGE_HEIGHT))
    for page in model.appendix_pages:
        pdf.setFillColorRGB(*[component / 255 for component in _hex_to_rgb(COLORS["background"])])
        pdf.rect(0, 0, PAGE_WIDTH, PAGE_HEIGHT, fill=1, stroke=0)
        for node in page.nodes:
            _draw_pdf_node(pdf, node, PAGE_HEIGHT)
        for block in page.blocks:
            _draw_pdf_text(pdf, block, PAGE_HEIGHT)
        pdf.showPage()
    pdf.save()


def _render_one_pager_pdf(model: ShowcaseModel, output_path: Path) -> None:
    page_width, page_height = landscape(A4)
    pdf = _new_pdf(output_path, (page_width, page_height))
    pdf.setFillColorRGB(*[component / 255 for component in _hex_to_rgb(COLORS["background"])])
    pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)
    for edge in model.one_pager.edges:
        _draw_pdf_edge(pdf, edge, page_height)
    for node in model.one_pager.nodes:
        _draw_pdf_node(pdf, node, page_height)
    for block in model.one_pager.blocks:
        _draw_pdf_text(pdf, block, page_height)
    pdf.showPage()
    pdf.save()


def _font_for_preview(size: int, bold: bool) -> ImageFont.ImageFont:
    candidates = (
        Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")
        if bold
        else Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
        if bold
        else Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _draw_preview_text(
    draw: ImageDraw.ImageDraw,
    block: TextBlock,
    scale_x: float,
    scale_y: float,
) -> None:
    font = _font_for_preview(max(10, round(block.font_size * scale_y)), block.bold)
    lines = block.text.splitlines() or [""]
    line_height = block.font_size * 1.22 * scale_y
    total_height = len(lines) * line_height
    y = block.y * scale_y + (block.height * scale_y - total_height) / 2
    for line in lines:
        box = draw.textbbox((0, 0), line, font=font)
        text_width = box[2] - box[0]
        x = block.x * scale_x
        if block.align == "center":
            x += (block.width * scale_x - text_width) / 2
        elif block.align == "right":
            x += block.width * scale_x - text_width
        draw.text((x, y), line, font=font, fill=_hex_to_rgb(block.color))
        y += line_height


def _render_preview_scene(scene: SlideScene, output_path: Path) -> None:
    scale_x = PREVIEW_WIDTH / PAGE_WIDTH
    scale_y = PREVIEW_HEIGHT / PAGE_HEIGHT
    image = Image.new("RGB", (PREVIEW_WIDTH, PREVIEW_HEIGHT), _hex_to_rgb(COLORS["background"]))
    draw = ImageDraw.Draw(image)
    for edge in scene.edges:
        draw.line(
            (
                edge.source_x * scale_x,
                edge.source_y * scale_y,
                edge.target_x * scale_x,
                edge.target_y * scale_y,
            ),
            fill=_hex_to_rgb(edge.color),
            width=3,
        )
    for node in scene.nodes:
        box = (
            node.x * scale_x,
            node.y * scale_y,
            (node.x + node.width) * scale_x,
            (node.y + node.height) * scale_y,
        )
        if node.shape == "circle":
            draw.ellipse(box, fill=_hex_to_rgb(node.fill), outline=_hex_to_rgb(node.stroke), width=2)
        elif node.shape == "rect":
            draw.rectangle(box, fill=_hex_to_rgb(node.fill), outline=_hex_to_rgb(node.stroke), width=2)
        else:
            draw.rounded_rectangle(box, radius=12, fill=_hex_to_rgb(node.fill), outline=_hex_to_rgb(node.stroke), width=2)
        _draw_preview_text(
            draw,
            TextBlock(
                node.x + 9,
                node.y + 6,
                node.width - 18,
                node.height - 12,
                node.label,
                node.font_size,
                node.text_color,
                node.bold,
                "center",
                "node",
            ),
            scale_x,
            scale_y,
        )
    for block in scene.blocks:
        _draw_preview_text(draw, block, scale_x, scale_y)
    image.save(output_path, format="PNG", optimize=False)


def render_validation_previews_v01(model: ShowcaseModel, preview_dir: str | Path) -> Path:
    target = Path(preview_dir)
    target.mkdir(mode=0o700)
    preview_paths: list[Path] = []
    for scene in model.slides:
        path = target / f"slide_{scene.slide_number:02d}.png"
        _render_preview_scene(scene, path)
        preview_paths.append(path)
    sheet = Image.new(
        "RGB",
        (PREVIEW_WIDTH * 4, PREVIEW_HEIGHT * 4),
        _hex_to_rgb(COLORS["background"]),
    )
    for index, path in enumerate(preview_paths):
        with Image.open(path) as image:
            sheet.paste(image, ((index % 4) * PREVIEW_WIDTH, (index // 4) * PREVIEW_HEIGHT))
    sheet_path = target / "contact_sheet.png"
    sheet.save(sheet_path, format="PNG", optimize=False)
    return sheet_path


def _claim_plain(claim: Claim) -> dict[str, object]:
    return {
        "claim_id": claim.claim_id,
        "display_label": claim.display_label,
        "observed_value": claim.observed_value,
        "observed_type": claim.observed_type,
        "evidence_availability": claim.evidence_availability,
        "source_path": claim.source_path,
        "source_field_or_section": claim.source_field_or_section,
        "source_sha256": claim.source_sha256,
        "source_commit": claim.source_commit,
        "slide_numbers": list(claim.slide_numbers),
        "appendix_pages": list(claim.appendix_pages),
        "public_link_available": claim.public_link_available,
        "notes": claim.notes,
    }


def _write_claim_matrix(model: ShowcaseModel, output_path: Path) -> None:
    raw = (
        json.dumps(
            [_claim_plain(claim) for claim in model.claims],
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
            allow_nan=False,
        ).encode("utf-8")
        + b"\n"
    )
    output_path.write_bytes(raw)


def _readme_text() -> str:
    return """# Hedgehog OS Airline All-Real Evidence Showcase v0.1

This static package presents accepted, committed proof-level Airline evidence.
It creates no new runtime evidence and makes no production claim.

## Deliverables

| File | Purpose |
| --- | --- |
| `README.md` | Package guide and regeneration instructions |
| `hedgehog_os_airline_all_real_showcase_v01.pptx` | Sixteen-slide main deck |
| `hedgehog_os_airline_all_real_showcase_v01.pdf` | Sixteen-page main deck PDF |
| `executive_one_pager_v01.pdf` | One-page executive summary |
| `technical_appendix_v01.pdf` | Nineteen-page technical appendix |
| `claim_evidence_matrix_v01.json` | Twenty-four claim-to-evidence bindings |
| `SHA256SUMS` | Checksums for the other six deliverables |

## Evidence Basis

- `docs/airline_all_real_full_stack_human_story_v01.md`
- `docs/airline_all_real_full_stack_crypto_anchor_v01.json`
- `docs/airline_all_real_full_stack_replay_report_v01.json`
- `docs/audit_reports/auditor_airline_all_real_full_stack_v01_generation_anchor_publication.log`
- `docs/audit_reports/auditor_airline_all_real_full_stack_v01_anchored_replay.log`
- `docs/airline_all_real_evidence_showcase_preflight_v01.md`

Committed files are linked directly in the claim matrix. Local sealed evidence
remains outside Git and is represented only through committed audits. The
renderer did not read a `.tmp` package. Raw prompts and raw responses are not embedded.

## Deck Outline

The sixteen slides cover the thesis, executive metrics, mock task, architecture,
twelve actors, Offer A selection, three Roots, five-phase Corridor, 19 / 29 / 3
Ledger, 9 / 11 Crypto generation, committed Anchor, Replay PASS, honest
FAIL_CLOSED history, recovery, proof claims, and non-claims.

The nineteen-page appendix contains evidence identity, twelve actor cards, the
full timeline, Crypto hashes, Replay zero-rerun boundaries, claim references,
and verification instructions.

## Regeneration

```sh
PYTHONPATH=. .venv/bin/python -m \\
  demo.run_airline_all_real_evidence_showcase_v01 \\
  --repository-root . \\
  --output-dir \\
  docs/showcase/airline_all_real_full_stack_v01
```

Regeneration performs no Gemini, provider, network, transaction, Corridor,
Ledger, Crypto, or Replay call.

Toolchain: Python 3.14.5; python-pptx 1.0.2; ReportLab 5.0.0; Pillow 12.3.0;
pypdf 6.14.2.

## Verify

macOS:

```sh
shasum -a 256 -c SHA256SUMS
```

GNU:

```sh
sha256sum -c SHA256SUMS
```

## Non-Claims

This is mock proof-level evidence: not a real ticket, booking, payment,
airline/bank/GDS integration, production deployment, signer authentication,
PKI, Root Attestation, semantic-truth proof, or effect authorization.
"""


def _write_readme(output_path: Path) -> None:
    output_path.write_text(_readme_text(), encoding="utf-8", newline="\n")


def _write_sha256sums(output_dir: Path) -> None:
    lines = []
    for filename in OUTPUT_NAMES:
        digest = hashlib.sha256((output_dir / filename).read_bytes()).hexdigest()
        lines.append(f"{digest}  {filename}")
    (output_dir / "SHA256SUMS").write_text("\n".join(lines) + "\n", encoding="ascii", newline="\n")


def _validate_generated_files(output_dir: Path) -> None:
    for filename in ALL_DELIVERABLE_NAMES:
        path = output_dir / filename
        path_stat = path.lstat()
        if stat.S_ISLNK(path_stat.st_mode) or not stat.S_ISREG(path_stat.st_mode):
            raise ValueError("deliverable_type")
        if path_stat.st_size == 0:
            raise ValueError("deliverable_empty")
    matrix = _strict_json((output_dir / "claim_evidence_matrix_v01.json").read_bytes())
    if not isinstance(matrix, list) or len(matrix) != 24:
        raise ValueError("matrix")
    with zipfile.ZipFile(output_dir / "hedgehog_os_airline_all_real_showcase_v01.pptx") as archive:
        slide_entries = [
            name
            for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        ]
        if len(slide_entries) != 16:
            raise ValueError("pptx_slides")
        if any(name.startswith("ppt/media/") for name in archive.namelist()):
            raise ValueError("pptx_media")
    from pypdf import PdfReader

    if len(PdfReader(output_dir / "hedgehog_os_airline_all_real_showcase_v01.pdf").pages) != 16:
        raise ValueError("deck_pages")
    if len(PdfReader(output_dir / "executive_one_pager_v01.pdf").pages) != 1:
        raise ValueError("one_pages")
    if len(PdfReader(output_dir / "technical_appendix_v01.pdf").pages) != 19:
        raise ValueError("appendix_pages")
    checksum_lines = (output_dir / "SHA256SUMS").read_text(encoding="ascii").splitlines()
    if len(checksum_lines) != 6:
        raise ValueError("checksums")
    for filename, line in zip(OUTPUT_NAMES, checksum_lines, strict=True):
        expected = hashlib.sha256((output_dir / filename).read_bytes()).hexdigest()
        if line != f"{expected}  {filename}":
            raise ValueError("checksum")


def _validate_path_input(value: str | Path) -> Path:
    if not isinstance(value, (str, Path)):
        raise ValueError("path_type")
    if isinstance(value, str) and (not value or "\x00" in value):
        raise ValueError("path_value")
    return Path(value)


def _require_no_symlink_components(path: Path) -> None:
    candidate = path if path.is_absolute() else Path.cwd() / path
    parts = candidate.parts
    current = Path(parts[0]) if candidate.is_absolute() else Path()
    for part in parts[1:] if candidate.is_absolute() else parts:
        current = current / part
        if current.exists() or current.is_symlink():
            if current.is_symlink():
                raise ValueError("symlink")


def generate_showcase_v01(
    *,
    repository_root: str | Path,
    output_dir: str | Path,
) -> dict[str, object]:
    owned_temp: Path | None = None
    created_parent: Path | None = None
    try:
        root = _validate_path_input(repository_root)
        output = _validate_path_input(output_dir)
        _require_no_symlink_components(root)
        _require_no_symlink_components(output.parent)
        root = root.resolve(strict=True)
        if not root.is_dir() or root.is_symlink():
            raise ValueError("root")
        if output.exists() or output.is_symlink():
            raise ValueError("output_exists")
        if not output.parent.exists():
            output_parent_parent = output.parent.parent.resolve(strict=True)
            if not output_parent_parent.is_dir() or output_parent_parent.is_symlink():
                raise ValueError("output_grandparent")
            output.parent.mkdir(mode=0o755)
            created_parent = output.parent.resolve(strict=True)
        parent = output.parent.resolve(strict=True)
        if not parent.is_dir() or parent.is_symlink():
            raise ValueError("output_parent")
        output_absolute = parent / output.name
        if ".tmp" in output_absolute.parts:
            raise ValueError("output_package")
        for source_path in SOURCE_COMMITS:
            source = (root / source_path).resolve(strict=True)
            if output_absolute == source or output_absolute in source.parents:
                raise ValueError("output_source")
        model = build_showcase_model_v01(root)
        owned_temp = Path(
            tempfile.mkdtemp(
                prefix=f".{output.name}.owned-",
                dir=parent,
            )
        )
        _write_readme(owned_temp / "README.md")
        _write_claim_matrix(model, owned_temp / "claim_evidence_matrix_v01.json")
        _render_pptx(
            model,
            owned_temp / "hedgehog_os_airline_all_real_showcase_v01.pptx",
            owned_temp,
        )
        _render_deck_pdf(
            model,
            owned_temp / "hedgehog_os_airline_all_real_showcase_v01.pdf",
        )
        _render_one_pager_pdf(
            model,
            owned_temp / "executive_one_pager_v01.pdf",
        )
        _render_appendix_pdf(
            model,
            owned_temp / "technical_appendix_v01.pdf",
        )
        _write_sha256sums(owned_temp)
        _validate_generated_files(owned_temp)
        if output_absolute.exists() or output_absolute.is_symlink():
            raise ValueError("output_race")
        os.rename(owned_temp, output_absolute)
        owned_temp = None
        created_parent = None
        return {
            "appendix_pages": 19,
            "deck_pdf_pages": 16,
            "deliverable_count": 7,
            "main_deck_slides": 16,
            "one_pager_pages": 1,
            "showcase_status": SHOWCASE_STATUS_PASS,
        }
    except Exception:
        if owned_temp is not None and owned_temp.exists():
            shutil.rmtree(owned_temp)
        if created_parent is not None:
            try:
                created_parent.rmdir()
            except OSError:
                pass
        raise _stable_error() from None


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog=MODULE_ID)
    parser.add_argument("--repository-root", required=True)
    parser.add_argument("--output-dir", required=True)
    try:
        args = parser.parse_args(argv)
        result = generate_showcase_v01(
            repository_root=args.repository_root,
            output_dir=args.output_dir,
        )
    except (ValueError, OSError):
        print(FAILURE_REASON)
        return 1
    print(json.dumps(result, sort_keys=True, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
