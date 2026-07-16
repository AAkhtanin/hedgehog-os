from __future__ import annotations

import ast
import hashlib
import json
import re
import shutil
import stat
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation

import demo.run_airline_all_real_evidence_showcase_v01 as showcase


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
RENDERER_PATH = REPOSITORY_ROOT / "demo/run_airline_all_real_evidence_showcase_v01.py"
SOURCE_PATHS = (
    showcase.PREFLIGHT_PATH,
    showcase.HUMAN_STORY_PATH,
    showcase.ANCHOR_PATH,
    showcase.REPLAY_REPORT_PATH,
    showcase.GENERATION_AUDIT_PATH,
    showcase.REPLAY_AUDIT_PATH,
)
EXPECTED_FILES = set(showcase.ALL_DELIVERABLE_NAMES)
EXPECTED_CLAIM_FIELDS = {
    "claim_id",
    "display_label",
    "observed_value",
    "observed_type",
    "evidence_availability",
    "source_path",
    "source_field_or_section",
    "source_sha256",
    "source_commit",
    "slide_numbers",
    "appendix_pages",
    "public_link_available",
    "notes",
}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _source_hashes(root: Path = REPOSITORY_ROOT) -> dict[str, str]:
    return {path: _sha256(root / path) for path in SOURCE_PATHS}


def _normalize(value: str) -> str:
    return " ".join(value.replace("\u2014", "-").replace("\u2013", "-").split())


def _pptx_text(path: Path) -> str:
    presentation = Presentation(path)
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text_frame")
    )


def _pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _copy_sources(target_root: Path) -> None:
    for relative_path in SOURCE_PATHS:
        destination = target_root / relative_path
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(REPOSITORY_ROOT / relative_path, destination)


@pytest.fixture(autouse=True)
def evidence_stability():
    before = _source_hashes()
    yield
    assert _source_hashes() == before


@pytest.fixture(scope="session")
def model():
    return showcase.build_showcase_model_v01(REPOSITORY_ROOT)


@pytest.fixture(scope="session")
def rendered_dir(tmp_path_factory):
    output = tmp_path_factory.mktemp("airline-showcase-render") / "showcase"
    result = showcase.generate_showcase_v01(
        repository_root=REPOSITORY_ROOT,
        output_dir=output,
    )
    assert result == {
        "appendix_pages": 19,
        "deck_pdf_pages": 16,
        "deliverable_count": 7,
        "main_deck_slides": 16,
        "one_pager_pages": 1,
        "showcase_status": "PASS",
    }
    return output


def test_static_boundary_uses_only_approved_inputs_and_dependencies():
    source = RENDERER_PATH.read_text(encoding="utf-8")
    tree = ast.parse(source)
    imported_roots = {
        alias.name.split(".")[0]
        for node in ast.walk(tree)
        if isinstance(node, ast.Import)
        for alias in node.names
    } | {
        (node.module or "").split(".")[0]
        for node in ast.walk(tree)
        if isinstance(node, ast.ImportFrom)
    }
    approved_external = {"PIL", "pptx", "reportlab", "pypdf"}
    external = imported_roots - {
        "__future__",
        "argparse",
        "hashlib",
        "json",
        "os",
        "re",
        "shutil",
        "stat",
        "tempfile",
        "zipfile",
        "dataclasses",
        "datetime",
        "pathlib",
        "typing",
    }
    assert external == approved_external
    forbidden = (
        "requests",
        "urllib",
        "socket",
        "config.py",
        "Path.glob",
        "Path.rglob",
        ".glob(",
        ".rglob(",
        "os.walk",
        "listdir(",
        "iterdir(",
        "getenv(",
        "environ[",
        "*_prompt.txt",
        "*_raw_response.txt",
    )
    assert not any(marker in source for marker in forbidden)
    assert tuple(showcase.SOURCE_COMMITS) == SOURCE_PATHS
    assert not any(path.startswith(".tmp/") for path in SOURCE_PATHS)


def test_source_validation_and_normalized_model(model):
    assert len(model.evidence_sources) == 6
    for source in model.evidence_sources:
        path = REPOSITORY_ROOT / source.path
        mode = path.lstat().st_mode
        assert stat.S_ISREG(mode)
        assert not stat.S_ISLNK(mode)
        assert source.source_sha256 == _sha256(path)
    assert tuple(actor.actor_id for actor in model.actors) == showcase.ACTOR_IDS
    assert all("PASS" in actor.validation_result for actor in model.actors)
    assert all("accepted: true" in actor.validation_result for actor in model.actors)
    assert all("Advisory only" in actor.authority_boundary for actor in model.actors)
    assert len(model.timeline) == 19
    assert [row.replay_index for row in model.timeline] == list(range(19))
    assert sum(row.dependency_count for row in model.timeline) == 29
    assert sum(row.is_root_final for row in model.timeline) == 3
    assert len(model.claims) == 24
    assert len(model.slides) == 16
    assert len(model.appendix_pages) == 19
    slide_five_labels = tuple(node.label for node in model.slides[4].nodes)
    assert slide_five_labels == showcase.HUMAN_ACTOR_LABELS
    assert not any(actor_id in "\n".join(slide_five_labels) for actor_id in showcase.ACTOR_IDS)
    slide_eleven_text = "\n".join(block.text for block in model.slides[10].blocks)
    for identifier in showcase.TRANSACTION_EVIDENCE_IDS + showcase.ROOT_FINAL_IDS:
        assert identifier in slide_eleven_text
    slide_fourteen_text = "\n".join(node.label for node in model.slides[13].nodes)
    assert "SELF_CONSISTENT_UNANCHORED" in slide_fourteen_text
    slide_sixteen_text = "\n".join(block.text for block in model.slides[15].blocks)
    for path in showcase.SHOWCASE_ENTRY_PATHS:
        assert path in slide_sixteen_text


def test_exact_deliverables_are_regular_nonempty_files(rendered_dir):
    entries = {path.name for path in rendered_dir.iterdir()}
    assert entries == EXPECTED_FILES
    for filename in entries:
        path = rendered_dir / filename
        mode = path.lstat().st_mode
        assert stat.S_ISREG(mode)
        assert not stat.S_ISLNK(mode)
        assert path.stat().st_size > 0
    assert not any(path.suffix.lower() in {".png", ".jpg", ".jpeg", ".svg"} for path in rendered_dir.iterdir())


def test_pptx_structure_titles_footers_fonts_and_metadata(rendered_dir):
    path = rendered_dir / "hedgehog_os_airline_all_real_showcase_v01.pptx"
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert len([name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]) == 16
        assert not any(name.startswith("ppt/media/") for name in names)
        relationship_text = "\n".join(
            archive.read(name).decode("utf-8")
            for name in names
            if name.endswith(".rels")
        )
        assert 'TargetMode="External"' not in relationship_text
        assert 'show="0"' not in "\n".join(
            archive.read(name).decode("utf-8")
            for name in names
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
    presentation = Presentation(path)
    assert len(presentation.slides) == 16
    assert presentation.slide_width == int(showcase.Inches(showcase.PPTX_WIDTH_INCHES))
    assert presentation.slide_height == 6_858_000
    for number, (slide, title) in enumerate(zip(presentation.slides, showcase.SLIDE_TITLES, strict=True), 1):
        text = "\n".join(
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text_frame")
        )
        assert _normalize(title) in _normalize(text)
        assert f"Evidence: " in text
        assert re.search(rf"(^|\n){number}($|\n)", text)
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            shape_text = shape.text.strip()
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if not run.text:
                        continue
                    assert run.font.name == "Arial"
                    assert run.font.size is not None
                    minimum = 10 if (
                        shape_text.startswith("Evidence:")
                        or shape_text.startswith("Evidence entry points:")
                        or shape_text == str(number)
                        or shape_text == "HEDGEHOG OS"
                    ) else 18
                    assert run.font.size.pt >= minimum
    slide_five_text = "\n".join(
        shape.text
        for shape in presentation.slides[4].shapes
        if hasattr(shape, "text_frame")
    )
    for label in showcase.HUMAN_ACTOR_LABELS:
        assert slide_five_text.count(label) == 1
    assert not any(actor_id in slide_five_text for actor_id in showcase.ACTOR_IDS)
    assert not any(
        line.rstrip().endswith(("_ll", "_reviewe", "_archi", "_explaine"))
        for line in slide_five_text.splitlines()
    )

    slide_eleven_text = "\n".join(
        shape.text
        for shape in presentation.slides[10].shapes
        if hasattr(shape, "text_frame")
    )
    for identifier in showcase.TRANSACTION_EVIDENCE_IDS + showcase.ROOT_FINAL_IDS:
        assert identifier in slide_eleven_text
    for label in (
        "Hold packet",
        "Client purchase intent",
        "Bank authorization ref",
        "Ticket issue intent",
        "Mock ticket receipt",
        "Mock purchase receipt",
        "Root finals",
    ):
        assert label in slide_eleven_text
    assert not any(line.strip() == "01" for line in slide_eleven_text.splitlines())
    assert not any(line.rstrip().endswith(":0") for line in slide_eleven_text.splitlines())

    slide_fourteen_text = "\n".join(
        shape.text
        for shape in presentation.slides[13].shapes
        if hasattr(shape, "text_frame")
    )
    assert "SELF_CONSISTENT_UNANCHORED" in slide_fourteen_text
    assert not any(line.strip() == "D" for line in slide_fourteen_text.splitlines())

    slide_sixteen_text = "\n".join(
        shape.text
        for shape in presentation.slides[15].shapes
        if hasattr(shape, "text_frame")
    )
    for evidence_path in showcase.SHOWCASE_ENTRY_PATHS:
        assert evidence_path in slide_sixteen_text
    properties = presentation.core_properties
    assert properties.title == "Hedgehog OS Airline All-Real Evidence Showcase v0.1"
    assert properties.author == "Hedgehog OS"
    assert properties.last_modified_by == "Hedgehog OS"
    assert properties.created.year == 2000
    assert properties.modified.year == 2000
    assert properties.revision == 1


def test_pdf_page_counts_titles_sizes_and_nonblank_text(rendered_dir):
    deck = PdfReader(rendered_dir / "hedgehog_os_airline_all_real_showcase_v01.pdf")
    one_pager = PdfReader(rendered_dir / "executive_one_pager_v01.pdf")
    appendix = PdfReader(rendered_dir / "technical_appendix_v01.pdf")
    assert len(deck.pages) == 16
    assert len(one_pager.pages) == 1
    assert len(appendix.pages) == 19
    for page, title in zip(deck.pages, showcase.SLIDE_TITLES, strict=True):
        assert float(page.mediabox.width) == pytest.approx(960, abs=0.01)
        assert float(page.mediabox.height) == pytest.approx(540, abs=0.01)
        text = page.extract_text() or ""
        assert _normalize(title) in _normalize(text)
        assert len(text.strip()) > 80
    for page, title in zip(appendix.pages, showcase.APPENDIX_TITLES, strict=True):
        assert float(page.mediabox.width) == pytest.approx(960, abs=0.01)
        assert float(page.mediabox.height) == pytest.approx(540, abs=0.01)
        text = page.extract_text() or ""
        assert _normalize(title) in _normalize(text)
        assert len(text.strip()) > 80
    one = one_pager.pages[0]
    assert float(one.mediabox.width) == pytest.approx(841.8898, abs=0.1)
    assert float(one.mediabox.height) == pytest.approx(595.2756, abs=0.1)
    one_text = one.extract_text() or ""
    assert "12 / 3 / 19 / 29 / 9 / 11 / 19 / 0" in one_text
    assert "No real ticket, booking, payment" in one_text
    for evidence_path in showcase.SHOWCASE_ENTRY_PATHS:
        assert evidence_path in one_text

    slide_five_text = deck.pages[4].extract_text() or ""
    for label in showcase.HUMAN_ACTOR_LABELS:
        assert slide_five_text.count(label) == 1
    assert not any(actor_id in slide_five_text for actor_id in showcase.ACTOR_IDS)
    assert not any(
        line.rstrip().endswith(("_ll", "_reviewe", "_archi", "_explaine"))
        for line in slide_five_text.splitlines()
    )

    slide_eleven_text = deck.pages[10].extract_text() or ""
    for identifier in showcase.TRANSACTION_EVIDENCE_IDS + showcase.ROOT_FINAL_IDS:
        assert identifier in slide_eleven_text
    assert not any(line.strip() == "01" for line in slide_eleven_text.splitlines())
    assert not any(line.rstrip().endswith(":0") for line in slide_eleven_text.splitlines())

    slide_fourteen_text = deck.pages[13].extract_text() or ""
    assert "SELF_CONSISTENT_UNANCHORED" in slide_fourteen_text
    assert not any(line.strip() == "D" for line in slide_fourteen_text.splitlines())

    slide_sixteen_text = deck.pages[15].extract_text() or ""
    for evidence_path in showcase.SHOWCASE_ENTRY_PATHS:
        assert evidence_path in slide_sixteen_text


def test_claim_matrix_exact_contract(rendered_dir):
    matrix = json.loads((rendered_dir / "claim_evidence_matrix_v01.json").read_text(encoding="utf-8"))
    assert isinstance(matrix, list)
    assert len(matrix) == 24
    assert [item["claim_id"] for item in matrix] == [f"C{number:02d}" for number in range(1, 25)]
    assert all(set(item) == EXPECTED_CLAIM_FIELDS for item in matrix)
    assert {number for item in matrix for number in item["slide_numbers"]} == set(range(1, 17))
    assert {number for item in matrix for number in item["appendix_pages"]} == set(range(1, 20))
    for item in matrix:
        assert item["source_path"] in SOURCE_PATHS
        assert (REPOSITORY_ROOT / item["source_path"]).is_file()
        assert re.fullmatch(r"[0-9a-f]{64}", item["source_sha256"])
        assert item["source_sha256"] == _sha256(REPOSITORY_ROOT / item["source_path"])
        assert item["source_commit"] == showcase.SOURCE_COMMITS[item["source_path"]]
        assert item["evidence_availability"] in {
            "committed_repository_evidence",
            "committed_audit_of_local_package",
        }
        assert item["public_link_available"] is True
        assert not item["source_path"].startswith(".tmp/")
        assert 1 <= min(item["slide_numbers"]) <= max(item["slide_numbers"]) <= 16
        assert 1 <= min(item["appendix_pages"]) <= max(item["appendix_pages"]) <= 19
    values = {item["claim_id"]: item["observed_value"] for item in matrix}
    assert values == {
        "C01": "PASS",
        "C02": 12,
        "C03": "12 PASS / 0 FAIL",
        "C04": "tri_airline_purchase:PAR-LIM:2026-08-12:client_001 / PAR to LIM / 2026-08-12",
        "C05": "offer:mock_airline_al:PAR-LIM:001",
        "C06": "4 PASS",
        "C07": "ClientRoot / AirlineRoot / BankRoot",
        "C08": "PASS",
        "C09": 1,
        "C10": 19,
        "C11": 29,
        "C12": 3,
        "C13": 9,
        "C14": 11,
        "C15": "SELF_CONSISTENT_UNANCHORED",
        "C16": True,
        "C17": "PASS",
        "C18": 19,
        "C19": "provider 0 / network 0 / Gemini 0",
        "C20": 0,
        "C21": True,
        "C22": True,
        "C23": "failed package promoted false; recovery package distinct true",
        "C24": False,
    }


def test_generated_readme_contract(rendered_dir):
    text = (rendered_dir / "README.md").read_text(encoding="utf-8")
    for filename in showcase.ALL_DELIVERABLE_NAMES:
        assert filename in text
    for source_path in SOURCE_PATHS:
        assert source_path in text
    assert "PYTHONPATH=. .venv/bin/python -m \\" in text
    assert "shasum -a 256 -c SHA256SUMS" in text
    assert "sha256sum -c SHA256SUMS" in text
    assert "renderer did not read a `.tmp` package" in text
    assert "Raw prompts and raw responses are not embedded" in text
    assert "performs no Gemini, provider, network" in text
    assert "not a real ticket" in text


def test_sha256sums_exact_order_and_values(rendered_dir):
    lines = (rendered_dir / "SHA256SUMS").read_text(encoding="ascii").splitlines()
    assert len(lines) == 6
    assert [line.split("  ", 1)[1] for line in lines] == sorted(showcase.OUTPUT_NAMES)
    for line in lines:
        digest, filename = line.split("  ", 1)
        assert re.fullmatch(r"[0-9a-f]{64}", digest)
        assert digest == _sha256(rendered_dir / filename)
        assert filename != "SHA256SUMS"


def test_security_and_nonclaim_wording(rendered_dir):
    text = "\n".join(
        (
            _pptx_text(rendered_dir / "hedgehog_os_airline_all_real_showcase_v01.pptx"),
            _pdf_text(rendered_dir / "hedgehog_os_airline_all_real_showcase_v01.pdf"),
            _pdf_text(rendered_dir / "executive_one_pager_v01.pdf"),
            _pdf_text(rendered_dir / "technical_appendix_v01.pdf"),
            (rendered_dir / "README.md").read_text(encoding="utf-8"),
            (rendered_dir / "claim_evidence_matrix_v01.json").read_text(encoding="utf-8"),
        )
    )
    forbidden = (
        "/Users/",
        "BEGIN PRIVATE KEY",
        "AIza",
        "production ready",
        "production-ready",
        "real ticket issued",
        "real booking created",
        "real payment executed",
        "signer authenticated",
        "PKI implemented",
        "Root Attestation implemented",
        "Exact Gemini quote",
        "raw provider response:",
        "raw prompt:",
    )
    lowered = text.lower()
    assert not any(marker.lower() in lowered for marker in forbidden)
    assert "NOT A REAL TICKET" in text
    assert "NOT A REAL BOOKING" in text
    assert "NOT A REAL PAYMENT" in text
    assert "signature verified false" in lowered or "signature unverified" in lowered


def test_deterministic_rendering_is_byte_identical(tmp_path):
    first = tmp_path / "first"
    second = tmp_path / "second"
    showcase.generate_showcase_v01(repository_root=REPOSITORY_ROOT, output_dir=first)
    showcase.generate_showcase_v01(repository_root=REPOSITORY_ROOT, output_dir=second)
    assert {path.name for path in first.iterdir()} == EXPECTED_FILES
    assert {path.name for path in second.iterdir()} == EXPECTED_FILES
    for filename in EXPECTED_FILES:
        assert (first / filename).read_bytes() == (second / filename).read_bytes()


def test_output_isolation_and_existing_output_fail_closed(tmp_path):
    first = tmp_path / "one"
    second = tmp_path / "two"
    showcase.generate_showcase_v01(repository_root=REPOSITORY_ROOT, output_dir=first)
    showcase.generate_showcase_v01(repository_root=REPOSITORY_ROOT, output_dir=second)
    (first / "user-note.txt").write_text("unrelated", encoding="utf-8")
    with pytest.raises(ValueError, match=showcase.FAILURE_REASON):
        showcase.generate_showcase_v01(repository_root=REPOSITORY_ROOT, output_dir=first)
    assert (first / "user-note.txt").read_text(encoding="utf-8") == "unrelated"
    assert {path.name for path in second.iterdir()} == EXPECTED_FILES


def test_missing_and_inconsistent_evidence_fail_closed(tmp_path):
    missing_root = tmp_path / "missing"
    missing_root.mkdir()
    missing_output = tmp_path / "missing-output"
    with pytest.raises(ValueError, match=showcase.FAILURE_REASON):
        showcase.generate_showcase_v01(repository_root=missing_root, output_dir=missing_output)
    assert not missing_output.exists()

    inconsistent_root = tmp_path / "inconsistent"
    inconsistent_root.mkdir()
    _copy_sources(inconsistent_root)
    replay_path = inconsistent_root / showcase.REPLAY_REPORT_PATH
    replay = json.loads(replay_path.read_text(encoding="utf-8"))
    replay["timeline_row_count"] = 18
    replay_path.write_text(json.dumps(replay, sort_keys=True) + "\n", encoding="utf-8")
    inconsistent_output = tmp_path / "inconsistent-output"
    with pytest.raises(ValueError, match=showcase.FAILURE_REASON):
        showcase.generate_showcase_v01(
            repository_root=inconsistent_root,
            output_dir=inconsistent_output,
        )
    assert not inconsistent_output.exists()


def test_owned_partial_output_is_removed(monkeypatch, tmp_path):
    output = tmp_path / "partial-output"
    user_file = tmp_path / "user-file.txt"
    user_file.write_text("keep", encoding="utf-8")

    def fail_render(*_args, **_kwargs):
        raise OSError("focused injected failure")

    monkeypatch.setattr(showcase, "_render_deck_pdf", fail_render)
    with pytest.raises(ValueError, match=showcase.FAILURE_REASON):
        showcase.generate_showcase_v01(repository_root=REPOSITORY_ROOT, output_dir=output)
    assert not output.exists()
    assert not any(path.name.startswith(f".{output.name}.owned-") for path in tmp_path.iterdir())
    assert user_file.read_text(encoding="utf-8") == "keep"


def test_pillow_previews_are_nonblank_colored_and_cleanable(model, tmp_path):
    preview_dir = tmp_path / "previews"
    sheet_path = showcase.render_validation_previews_v01(model, preview_dir)
    preview_paths = sorted(preview_dir.glob("slide_*.png"))
    assert len(preview_paths) == 16
    background = showcase._hex_to_rgb(showcase.COLORS["background"])
    accents = {
        showcase._hex_to_rgb(showcase.COLORS[name])
        for name in ("cyan", "client", "airline", "bank", "pass", "advisory", "fail")
    }
    for path in preview_paths:
        with Image.open(path) as image:
            assert image.size == (1600, 900)
            colors = image.getcolors(maxcolors=2_000_000)
            assert colors is not None
            non_background = sum(count for count, color in colors if color != background)
            assert non_background > 15_000
            assert any(color in accents for _, color in colors)
    with Image.open(sheet_path) as sheet:
        assert sheet.size == (6400, 3600)
        assert sheet.getbbox() is not None
    shutil.rmtree(preview_dir)
    assert not preview_dir.exists()
